import shutil
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any, List, Dict, Tuple

# ============================================================
# 配置区 —— 日常使用主要修改这里
# 高速版：一个 Excel 只打开一次，并使用流式行读取。
# ============================================================

CONFIG = {
    # 当前脚本所在目录
    "input_dir": Path(__file__).resolve().parent,

    # 输出文件名
    "output_file": "03-嘉盛物料映射汇总表.xlsx",

    # 异常文件移动目录（读取失败的文件会移动到这里）
    "error_dir": "异常表",

    # 正常完成文件移动目录（处理成功的文件会移动到这里）
    "done_dir": "已读取",

    # 日志文件名
    "log_file": "02-处理日志.log",

    # 并行读取文件数（1 为串行，建议 4~8）
    "max_workers": 8,

    # 是否在最终结果后面增加来源信息
    "add_source_info": True,

    # 去重依据：
    # 最终统一结构中的第几列，从 1 开始
    # 当前第 1 列是序号，第 2 列是物料编码
    "dedup_col": 2,

    # 数量所在的目标列（最终结构中的第几列，从 1 开始）
    # G 列 = 第 7 列
    # 当该列写入的值不是数字时，自动从原表头查找列名为"数量"或"用量"的列进行替换
    "quantity_col": 7,

    # 物料映射列配置
    # 从哪一列开始放物料映射（物料A、物料A数量、物料B、物料B数量……）
    "mapping_start_col": "M",

    # 最多放几种物料映射
    "mapping_count": 5,

    # 每个 Sheet 独立配置
    "sheets": [
        {
            "name": "电柜元器件",

            # 从 Excel 第几行开始读取
            "start_row": 4,

            # 原始读取范围
            "start_col": "A",
            "end_col": "F",

            # 以哪一列为空作为数据结束
            "stop_col": "A",

            # 字段映射：
            # 最终统一结构一共 7 列
            #
            # 原始：
            # A B C D E F
            #
            # 最终：
            # A B C D E 空 F
            #
            # 数字表示原始区域中的第几列
            # None 表示插入空值
            "column_mapping": [
                1,  # A
                2,  # B
                3,  # C
                4,  # D
                5,  # E
                None,
                6,  # 原 F -> 最终 G
            ],
        },

        {
            "name": "电柜半成品",

            "start_row": 3,
            "start_col": "A",
            "end_col": "G",
            "stop_col": "A",

            # A:G 原样映射到最终 7 列
            "column_mapping": [
                1,
                2,
                3,
                4,
                5,
                6,
                7,
            ],
        },
    ],
}


# ============================================================
# 日志
# ============================================================

_log_lines: List[str] = []


def log(msg: str = ""):
    """同时输出到控制台和日志缓冲区。"""
    print(msg)
    _log_lines.append(msg)


def save_log(log_path: Path):
    log_path.write_text(
        "\n".join(_log_lines),
        encoding="utf-8",
    )


# ============================================================
# 工具函数
# ============================================================

def excel_col_to_num(col: str) -> int:
    """
    Excel 列字母转数字。
    A -> 1
    B -> 2
    Z -> 26
    AA -> 27
    """
    col = col.upper().strip()

    result = 0

    for char in col:
        result = result * 26 + ord(char) - ord("A") + 1

    return result


def is_empty(value: Any) -> bool:
    """
    判断单元格是否为空。
    """
    if value is None:
        return True

    if isinstance(value, str) and value.strip() == "":
        return True

    return False


def is_number(value: Any) -> bool:
    """
    判断值是否为数字。
    """
    if isinstance(value, (int, float)):
        return True

    if isinstance(value, str):
        s = value.strip()
        if s:
            try:
                float(s)
                return True
            except ValueError:
                return False

    return False


def normalize_key(value: Any) -> str:
    """
    将物料编码标准化成字符串用于去重。

    避免 Excel 中：
    123
    123.0
    "123"
    因类型不同而导致误判。
    """
    if value is None:
        return ""

    # Excel 数字经常会读成浮点数
    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def apply_mapping(
    source_row: List[Any],
    mapping: List[Any],
) -> List[Any]:
    """
    根据配置将原始数据转换成统一结构。

    mapping 示例：
    [1, 2, 3, 4, 5, None, 6]

    表示：

    输出1 <- 原始1
    输出2 <- 原始2
    ...
    输出6 <- 空
    输出7 <- 原始6
    """

    result = []

    for source_index in mapping:

        if source_index is None:
            result.append(None)
            continue

        index = source_index - 1

        if index < len(source_row):
            result.append(source_row[index])
        else:
            result.append(None)

    return result


def check_header_b(b_value: Any) -> bool:
    """
    检查 B 列表头是否为物料号/料号。
    """
    if is_empty(b_value):
        return False
    s = str(b_value).strip()
    return "料号" in s


def mapping_label(i: int) -> str:
    """
    生成物料映射标签：A, B, ..., Z, AA, AB, ...
    """
    result = ""
    n = i + 1
    while n > 0:
        n, rem = divmod(n - 1, 26)
        result = chr(ord("A") + rem) + result
    return result


# ============================================================
# Excel 高速读取
# ============================================================

def _read_xlsx_sheet_from_workbook(
    workbook,
    sheet_config: Dict[str, Any],
) -> List[List[Any]]:
    """
    从已经打开的 xlsx/xlsm 工作簿读取一个 Sheet。

    性能优化：
    1. 同一个 Excel 文件只打开一次；
    2. 使用 iter_rows(values_only=True) 流式读取，避免逐单元格 ws.cell()；
    3. 一次读取到本 Sheet 实际需要的最右列（含数量备用列搜索范围）。
    """
    sheet_name = sheet_config["name"]

    if sheet_name not in workbook.sheetnames:
        log(f"    [跳过] Sheet 不存在：{sheet_name}")
        return []

    ws = workbook[sheet_name]

    start_row = sheet_config["start_row"]
    header_row = start_row - 1

    start_col = excel_col_to_num(sheet_config["start_col"])
    end_col = excel_col_to_num(sheet_config["end_col"])
    stop_col = excel_col_to_num(sheet_config["stop_col"])

    # 一次读取表头前 30 列，用于检查 B 列及查找“数量/用量”
    header_values = next(
        ws.iter_rows(
            min_row=header_row,
            max_row=header_row,
            min_col=1,
            max_col=30,
            values_only=True,
        ),
        (),
    )

    b_header = header_values[1] if len(header_values) >= 2 else None
    if not check_header_b(b_header):
        raise ValueError(
            f"B 列表头不是物料号/料号，实际为：{b_header}"
        )

    quantity_col = None
    for idx, value in enumerate(header_values, start=1):
        if value and any(
            k in str(value).strip()
            for k in ("数量", "用量")
        ):
            quantity_col = idx
            break

    qty_index = CONFIG["quantity_col"] - 1
    rows: List[List[Any]] = []

    # 只流式读取真正需要的列：
    # 原始范围、停止列、以及可能作为数量备用来源的列。
    max_needed_col = max(
        end_col,
        stop_col,
        quantity_col or 0,
    )

    for values in ws.iter_rows(
        min_row=start_row,
        min_col=1,
        max_col=max_needed_col,
        values_only=True,
    ):
        stop_value = (
            values[stop_col - 1]
            if stop_col - 1 < len(values)
            else None
        )

        if is_empty(stop_value):
            break

        source_row = list(values[start_col - 1:end_col])

        mapped_row = apply_mapping(
            source_row,
            sheet_config["column_mapping"],
        )

        if (
            quantity_col is not None
            and not is_number(mapped_row[qty_index])
        ):
            mapped_row[qty_index] = (
                values[quantity_col - 1]
                if quantity_col - 1 < len(values)
                else None
            )

        rows.append(mapped_row)

    return rows


def _read_xls_sheet_from_workbook(
    workbook,
    sheet_config: Dict[str, Any],
) -> List[List[Any]]:
    """
    从已经打开的 .xls 工作簿读取一个 Sheet。
    同一个 .xls 文件也只打开一次。
    """
    import xlrd

    sheet_name = sheet_config["name"]

    try:
        ws = workbook.sheet_by_name(sheet_name)
    except xlrd.biffh.XLRDError:
        log(f"    [跳过] Sheet 不存在：{sheet_name}")
        return []

    start_row = sheet_config["start_row"] - 1
    start_col = excel_col_to_num(sheet_config["start_col"]) - 1
    end_col = excel_col_to_num(sheet_config["end_col"]) - 1
    stop_col = excel_col_to_num(sheet_config["stop_col"]) - 1

    header_row_0 = start_row - 1
    b_header = ws.cell_value(header_row_0, 1)
    if not check_header_b(b_header):
        raise ValueError(
            f"B 列表头不是物料号/料号，实际为：{b_header}"
        )

    quantity_col = None
    for col in range(min(ws.ncols, 30)):
        value = ws.cell_value(header_row_0, col)
        if value and any(
            k in str(value).strip()
            for k in ("数量", "用量")
        ):
            quantity_col = col
            break

    rows: List[List[Any]] = []
    qty_index = CONFIG["quantity_col"] - 1

    for current_row in range(start_row, ws.nrows):
        stop_value = (
            ws.cell_value(current_row, stop_col)
            if stop_col < ws.ncols
            else None
        )

        if is_empty(stop_value):
            break

        source_row = [
            ws.cell_value(current_row, col)
            if col < ws.ncols
            else None
            for col in range(start_col, end_col + 1)
        ]

        mapped_row = apply_mapping(
            source_row,
            sheet_config["column_mapping"],
        )

        if (
            quantity_col is not None
            and not is_number(mapped_row[qty_index])
            and quantity_col < ws.ncols
        ):
            mapped_row[qty_index] = ws.cell_value(
                current_row,
                quantity_col,
            )

        rows.append(mapped_row)

    return rows


# ============================================================
# 输出 Excel
# ============================================================

def build_header() -> List[str]:
    """
    构建标题行。
    A~I: 序号、物料号、品名、型号、品牌、规格、数量、文件名、sheet名
    J: 处理时间
    K~(mapping_start_col-1): 空
    mapping_start_col 起: 物料A、物料A数量、物料B、物料B数量……
    """
    header = [
        "序号", "物料号", "品名", "型号", "品牌",
        "规格", "数量", "文件名", "sheet名", "处理时间",
    ]

    # 填充空白列到 mapping_start_col 前一列
    start_num = excel_col_to_num(CONFIG["mapping_start_col"])
    while len(header) < start_num - 1:
        header.append("")

    # 物料映射列
    for i in range(CONFIG["mapping_count"]):
        label = mapping_label(i)
        header.append(f"物料{label}")
        header.append(f"物料{label}数量")

    return header


def read_existing_output(output_path: Path) -> Tuple[set, int]:
    """
    读取已有输出文件中的物料号集合和最大序号，
    用于追加模式去重和序号续接。
    """
    from openpyxl import load_workbook

    existing_keys = set()
    max_seq = 0

    if not output_path.exists():
        return existing_keys, max_seq

    try:
        wb = load_workbook(
            output_path,
            read_only=True,
            data_only=True,
        )
        ws = wb.active

        for row in ws.iter_rows(
            min_row=2,
            values_only=True,
        ):
            if row and len(row) > 1:
                if row[0] is not None:
                    try:
                        seq = int(row[0])
                        if seq > max_seq:
                            max_seq = seq
                    except (ValueError, TypeError):
                        pass
                if row[1] is not None:
                    key = normalize_key(row[1])
                    if key:
                        existing_keys.add(key)

        wb.close()
    except Exception:
        pass

    return existing_keys, max_seq


def write_output(
    output_path: Path,
    rows: List[List[Any]],
    header: List[str],
):

    from openpyxl import Workbook

    wb = Workbook()

    ws = wb.active
    ws.title = "汇总"

    ws.append(header)

    for row in rows:
        padded = list(row) + [None] * (
            len(header) - len(row)
        )
        ws.append(padded)

    wb.save(output_path)


def append_output(
    output_path: Path,
    rows: List[List[Any]],
    header: List[str],
):

    from openpyxl import load_workbook

    if not output_path.exists():
        write_output(output_path, rows, header)
        return

    wb = load_workbook(output_path)
    ws = wb.active

    for row in rows:
        padded = list(row) + [None] * (
            len(header) - len(row)
        )
        ws.append(padded)

    wb.save(output_path)


# ============================================================
# 单文件处理（供并行调用）
# ============================================================

def process_one_file(
    file_path: Path,
) -> Tuple[Path, bool, List[Dict[str, Any]], str]:
    """
    处理单个 Excel 文件，返回：
    (file_path, file_ok, sheet_results, error_msg)

    高速版关键点：
    - 一个文件只打开一次；
    - 在同一个已打开工作簿中连续读取所有配置 Sheet；
    - 文件级并行仍由 ThreadPoolExecutor 负责。
    """
    file_ok = True
    sheet_results: List[Dict[str, Any]] = []
    errors: List[str] = []

    suffix = file_path.suffix.lower()

    try:
        if suffix in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook

            workbook = load_workbook(
                file_path,
                read_only=True,
                data_only=True,
            )

            try:
                for sheet_config in CONFIG["sheets"]:
                    sheet_name = sheet_config["name"]

                    try:
                        rows = _read_xlsx_sheet_from_workbook(
                            workbook,
                            sheet_config,
                        )
                        if rows:
                            sheet_results.append({
                                "sheet_name": sheet_name,
                                "rows": rows,
                            })
                    except Exception as exc:
                        file_ok = False
                        errors.append(f"{sheet_name}: {exc}")
            finally:
                workbook.close()

        elif suffix == ".xls":
            try:
                import xlrd
            except ImportError:
                raise RuntimeError(
                    "检测到 .xls 文件，但没有安装 xlrd。\n"
                    "请执行：pip install xlrd"
                )

            workbook = xlrd.open_workbook(
                file_path,
                on_demand=True,
            )

            try:
                for sheet_config in CONFIG["sheets"]:
                    sheet_name = sheet_config["name"]

                    try:
                        rows = _read_xls_sheet_from_workbook(
                            workbook,
                            sheet_config,
                        )
                        if rows:
                            sheet_results.append({
                                "sheet_name": sheet_name,
                                "rows": rows,
                            })
                    except Exception as exc:
                        file_ok = False
                        errors.append(f"{sheet_name}: {exc}")
            finally:
                workbook.release_resources()

        else:
            raise ValueError(
                f"不支持的 Excel 类型：{suffix}"
            )

    except Exception as exc:
        file_ok = False
        errors.append(str(exc))

    return (
        file_path,
        file_ok,
        sheet_results,
        " | ".join(errors),
    )


# ============================================================
# 主程序
# ============================================================

def main():

    input_dir = CONFIG["input_dir"]

    output_path = (
        input_dir
        / CONFIG["output_file"]
    )

    error_dir = input_dir / CONFIG["error_dir"]

    done_dir = input_dir / CONFIG["done_dir"]

    log_path = input_dir / CONFIG["log_file"]

    batch_time = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    log("=" * 60)
    log("Excel 物料汇总工具")
    log(f"处理时间：{batch_time}")
    log("=" * 60)

    # --------------------------------------------------------
    # 找 Excel
    # --------------------------------------------------------

    excel_files = []

    for file_path in input_dir.iterdir():

        if not file_path.is_file():
            continue

        if file_path.name.startswith("~$"):
            continue

        if file_path.name == CONFIG["output_file"]:
            continue

        if file_path.name == CONFIG["log_file"]:
            continue

        if file_path.suffix.lower() in (
            ".xls",
            ".xlsx",
            ".xlsm",
        ):
            excel_files.append(file_path)

    excel_files.sort()

    log(
        f"\n扫描目录：{input_dir}"
    )

    log(
        f"发现 Excel 文件：{len(excel_files)} 个\n"
    )

    # --------------------------------------------------------
    # 读取已有输出（追加模式）
    # --------------------------------------------------------

    seen_keys, next_seq = read_existing_output(
        output_path
    )

    if seen_keys:
        log(
            f"已有物料记录：{len(seen_keys)} 条，"
            f"最大序号：{next_seq}\n"
        )

    # --------------------------------------------------------
    # 统计
    # --------------------------------------------------------

    file_success = 0
    file_failed = 0
    file_moved = 0
    file_done_moved = 0

    raw_count = 0
    duplicate_count = 0

    result_rows = []

    # --------------------------------------------------------
    # 并行扫描文件
    # --------------------------------------------------------

    max_workers = CONFIG["max_workers"]

    log(f"并行读取：{max_workers} 线程\n")

    all_results: List[Tuple[Path, bool, List[Dict[str, Any]], str]] = []

    with ThreadPoolExecutor(
        max_workers=max_workers,
    ) as executor:

        future_map = {
            executor.submit(
                process_one_file, fp
            ): fp
            for fp in excel_files
        }

        done_count = 0

        for future in as_completed(future_map):

            fp = future_map[future]
            done_count += 1

            try:
                result = future.result()
            except Exception as exc:
                result = (fp, False, [], str(exc))

            all_results.append(result)

            file_path_r, file_ok_r, sheet_results_r, error_msg_r = result

            log(
                f"[{done_count}/{len(excel_files)}] "
                f"{'OK' if file_ok_r else 'FAIL'} "
                f"{file_path_r.name}"
            )

            if file_ok_r:
                for sr in sheet_results_r:
                    log(
                        f"    {sr['sheet_name']}: "
                        f"{len(sr['rows'])} 行"
                    )
            else:
                log(f"    [错误] {error_msg_r}")

    # --------------------------------------------------------
    # 合并去重 + 移动文件（保持文件顺序）
    # --------------------------------------------------------

    file_order = {
        file_path: index
        for index, file_path in enumerate(excel_files)
    }

    all_results.sort(
        key=lambda r: file_order[r[0]]
    )

    for file_path, file_ok, sheet_results, error_msg in all_results:

        if file_ok:

            for sr in sheet_results:

                sheet_name = sr["sheet_name"]
                rows = sr["rows"]

                for row in rows:

                    raw_count += 1

                    dedup_index = (
                        CONFIG["dedup_col"] - 1
                    )

                    key = normalize_key(
                        row[dedup_index]
                    )

                    if not key:
                        continue

                    if key in seen_keys:
                        duplicate_count += 1
                        continue

                    seen_keys.add(key)

                    output_row = list(row)

                    if CONFIG["add_source_info"]:
                        output_row.extend([
                            file_path.name,
                            sheet_name,
                        ])

                    output_row.append(batch_time)

                    result_rows.append(output_row)

            file_success += 1

            try:
                done_dir.mkdir(exist_ok=True)
                dest = done_dir / file_path.name
                if dest.exists():
                    dest.unlink()
                shutil.move(str(file_path), str(dest))
                file_done_moved += 1
                log(
                    f"    -> 已移动到："
                    f"{CONFIG['done_dir']}/"
                    f"{file_path.name}"
                )
            except Exception as move_exc:
                log(f"    [移动失败] {move_exc}")

        else:

            file_failed += 1

            try:
                error_dir.mkdir(exist_ok=True)
                dest = error_dir / file_path.name
                if dest.exists():
                    dest.unlink()
                shutil.move(str(file_path), str(dest))
                file_moved += 1
                log(
                    f"    -> 已移动到："
                    f"{CONFIG['error_dir']}/"
                    f"{file_path.name}"
                )
            except Exception as move_exc:
                log(f"    [移动失败] {move_exc}")

    # --------------------------------------------------------
    # 输出
    # --------------------------------------------------------

    header = build_header()

    if result_rows:

        try:
            append_output(
                output_path,
                result_rows,
                header,
            )
        except PermissionError:
            log(
                f"[错误] 输出文件被占用，"
                f"请关闭 Excel 后重试：{output_path}"
            )
        except Exception as exc:
            log(
                f"[错误] 写入输出文件失败：{exc}"
            )

    # --------------------------------------------------------
    # 结果
    # --------------------------------------------------------

    log("\n" + "=" * 60)
    log("处理完成")
    log("=" * 60)

    log(
        f"Excel 文件总数：{len(excel_files)}"
    )

    log(
        f"完全成功文件：{file_success}"
    )

    log(
        f"存在异常文件：{file_failed}"
    )

    log(
        f"已移动异常文件：{file_moved}"
    )

    log(
        f"已移动已读取文件：{file_done_moved}"
    )

    log(
        f"原始物料条数：{raw_count}"
    )

    log(
        f"重复物料条数：{duplicate_count}"
    )

    log(
        f"本次新增物料条数：{len(result_rows)}"
    )

    log(
        f"\n输出文件：{output_path}"
    )

    log(
        f"日志文件：{log_path}"
    )

    save_log(log_path)

    input(
        "\n按回车键退出..."
    )


# ============================================================
# 程序入口
# ============================================================

if __name__ == "__main__":
    main()
