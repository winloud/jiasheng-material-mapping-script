from __future__ import annotations

import re
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"

TEMPLATE_FILE = PROJECT_ROOT / "templates" / "ERP物料清单批量导入表.xls"
OUTPUT_DIR = PROJECT_ROOT / "_runtime_unused"
ERROR_DIR = PROJECT_ROOT / "_runtime_unused"
FILLBACK_OUTPUT_DIR = PROJECT_ROOT / "_runtime_unused"
LOG_FILE = PROJECT_ROOT / "_runtime_unused" / "erp_bom_export.log"

SUMMARY_SHEET = "京能物料对应汇总"
ERP_HEADERS = ("母件物料编码", "子件物料编码", "数量", "顺序号")

NUMBER_TOKEN_RE = re.compile(
    r"^(?P<start>\d+)(?:~(?P<end>\d+))?(?:\([^)]*\))?$"
)
BASE_RE = re.compile(r"^[A-Za-z]+\d+$")

LOG_LINES: List[str] = []


@dataclass
class ParentGroup:
    parent_material: str
    source_row: int
    project_keys: List[str] = field(default_factory=list)
    st_children: List[str] = field(default_factory=list)


def log(message: str = "") -> None:
    print(message)
    LOG_LINES.append(message)


def save_log() -> None:
    LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    LOG_FILE.write_text("\n".join(LOG_LINES), encoding="utf-8")


def is_blank(value: Any) -> bool:
    return value is None or (
        isinstance(value, str) and value.strip() == ""
    )


def normalize(value: Any) -> str:
    if is_blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()


def unique_path(directory: Path, filename: str) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / filename
    if not target.exists():
        return target

    p = Path(filename)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return directory / f"{p.stem}_{stamp}{p.suffix}"


# ============================================================
# 文件名项目号展开
# ============================================================

def expand_filename_codes(filename_or_stem: str) -> List[str]:
    """
    示例：
    HZJF26071-26~29-47~50-55-64~67-72~75-电气清单...
    HZJF26071-76~83(有高效)-电气清单...
    """
    stem = Path(filename_or_stem).stem.strip()
    parts = stem.split("-")

    if len(parts) < 2:
        return []

    base = parts[0].strip()
    if not BASE_RE.fullmatch(base):
        return []

    result: List[str] = []

    for token in parts[1:]:
        token = token.strip()
        match = NUMBER_TOKEN_RE.fullmatch(token)

        if not match:
            break

        start_text = match.group("start")
        end_text = match.group("end")

        if end_text is None:
            result.append(f"{base}-{start_text}")
            continue

        start = int(start_text)
        end = int(end_text)

        if end < start:
            raise ValueError(f"反向编号范围：{token}")

        width = max(len(start_text), len(end_text))
        for number in range(start, end + 1):
            result.append(f"{base}-{number:0{width}d}")

    return result


def source_stem_from_fillback(path: Path) -> str:
    """
    脚本3输出：
      原文件名_已匹配.xlsx
      原文件名_已匹配_YYYYMMDD_HHMMSS.xlsx
    还原成原电气清单文件 stem。
    """
    stem = path.stem

    marker = "_已匹配"
    pos = stem.find(marker)
    if pos < 0:
        return ""

    return stem[:pos]


def build_fillback_code_index() -> Tuple[Dict[str, Path], List[str]]:
    """
    直接从脚本3输出建立：
        HZJF26071-1 -> 对应的 _已匹配.xlsx

    这样脚本4不需要再次复制原始电气清单文件夹。
    如果同一个源清单存在多次脚本3输出，自动取最新文件。
    """
    FILLBACK_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # 先按“原始清单 stem”归并重复运行结果，取最新。
    latest_by_source: Dict[str, Path] = {}

    for path in FILLBACK_OUTPUT_DIR.iterdir():
        if not path.is_file():
            continue
        if path.name.startswith("~$"):
            continue
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            continue

        source_stem = source_stem_from_fillback(path)
        if not source_stem:
            continue

        old = latest_by_source.get(source_stem)
        if old is None or path.stat().st_mtime > old.stat().st_mtime:
            latest_by_source[source_stem] = path

    by_code: Dict[str, Path] = {}
    source_for_code: Dict[str, str] = {}
    errors: List[str] = []

    for source_stem, fillback in sorted(latest_by_source.items()):
        try:
            codes = expand_filename_codes(source_stem)
        except Exception as exc:
            errors.append(f"{source_stem}：{exc}")
            continue

        if not codes:
            # 输出目录可能存在非电气清单文件，不作为全局错误。
            continue

        for code in codes:
            key = code.upper()

            old_source = source_for_code.get(key)
            if old_source is not None and old_source != source_stem:
                errors.append(
                    f"{code} 同时对应两个不同清单："
                    f"{old_source} / {source_stem}"
                )
                continue

            by_code[key] = fillback
            source_for_code[key] = source_stem

    return by_code, errors


# ============================================================
# 物料导入.xls
# ============================================================

def find_material_import_file(project_dir: Path) -> Path:
    matches = [
        p for p in project_dir.iterdir()
        if p.is_file()
        and p.suffix.lower() == ".xls"
        and p.name.startswith("物料导入")
        and not p.name.startswith("~$")
    ]

    if len(matches) != 1:
        raise RuntimeError(
            f"项目目录需要且只能有一个“物料导入*.xls”，"
            f"当前找到 {len(matches)} 个。"
        )

    return matches[0]


def read_material_import(path: Path) -> List[ParentGroup]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError(
            "缺少 xlrd。请通过启动工具运行，程序会自动安装依赖。"
        ) from exc

    wb = xlrd.open_workbook(str(path), on_demand=True)

    try:
        ws = wb.sheet_by_index(0)
        groups: List[ParentGroup] = []
        current: Optional[ParentGroup] = None

        for row_index in range(ws.nrows):
            a = normalize(ws.cell_value(row_index, 0)) if ws.ncols >= 1 else ""
            b = normalize(ws.cell_value(row_index, 1)) if ws.ncols >= 2 else ""

            if not a:
                continue

            code_upper = a.upper()

            if code_upper.startswith("ET"):
                current = ParentGroup(
                    parent_material=a,
                    source_row=row_index + 1,
                )
                if b:
                    current.project_keys.append(b)
                groups.append(current)
                continue

            if code_upper.startswith("ST"):
                if current is None:
                    raise RuntimeError(
                        f"第 {row_index + 1} 行出现 ST，"
                        f"但前面没有 ET 母件：{a}"
                    )

                current.st_children.append(a)
                if b:
                    current.project_keys.append(b)

        if not groups:
            raise RuntimeError("A列没有找到 ET 开头的母件。")

        return groups
    finally:
        wb.release_resources()


def resolve_group_project_key(group: ParentGroup) -> str:
    unique: List[str] = []
    seen = set()

    for raw in group.project_keys:
        key = normalize(raw)
        upper = key.upper()
        if key and upper not in seen:
            seen.add(upper)
            unique.append(key)

    if not unique:
        raise RuntimeError(
            f"母件 {group.parent_material} 没有可用的 B 列项目编号。"
        )

    if len(unique) > 1:
        raise RuntimeError(
            f"母件 {group.parent_material} 对应多个不同B列项目编号："
            + " / ".join(unique)
        )

    return unique[0]


def quantity_to_decimal(value: Any, material: str) -> Decimal:
    text = normalize(value)
    if not text:
        raise RuntimeError(f"子件 {material} 的数量为空。")
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise RuntimeError(f"子件 {material} 的数量不是有效数字：{text}") from exc
    if not number.is_finite():
        raise RuntimeError(f"子件 {material} 的数量不是有限数字：{text}")
    return number


def decimal_to_excel_value(value: Decimal) -> Any:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


# ============================================================
# 京能物料对应汇总
# ============================================================

def read_jingneng_children(fillback_file: Path) -> List[Tuple[str, Any]]:
    wb = load_workbook(
        fillback_file,
        read_only=True,
        data_only=True,
    )

    try:
        if SUMMARY_SHEET not in wb.sheetnames:
            raise RuntimeError(
                f"{fillback_file.name} 缺少 Sheet：{SUMMARY_SHEET}"
            )

        ws = wb[SUMMARY_SHEET]
        rows = ws.iter_rows(values_only=True)

        try:
            header = next(rows)
        except StopIteration:
            raise RuntimeError(
                f"{fillback_file.name} 的 {SUMMARY_SHEET} 是空表。"
            )

        header_index = {
            normalize(name): index
            for index, name in enumerate(header)
            if not is_blank(name)
        }

        required = ("京能物料号", "京能数量", "匹配状态")
        missing = [x for x in required if x not in header_index]

        if missing:
            raise RuntimeError(
                f"缺少列：{', '.join(missing)}"
            )

        material_idx = header_index["京能物料号"]
        qty_idx = header_index["京能数量"]
        status_idx = header_index["匹配状态"]

        result: List[Tuple[str, Any]] = []

        for values in rows:
            status = normalize(
                values[status_idx] if status_idx < len(values) else None
            )

            if status != "已匹配":
                continue

            material = normalize(
                values[material_idx] if material_idx < len(values) else None
            )
            if not material:
                continue

            # “京能物料号”列直接标记为“客供件”的记录不进入ERP BOM。
            # 先于数量检查过滤，客供件即使数量为空也不会造成脚本4异常。
            if material == "客供件":
                continue

            qty = values[qty_idx] if qty_idx < len(values) else None
            if is_blank(qty):
                raise RuntimeError(
                    f"京能物料 {material} 的京能数量为空。"
                )

            result.append((material, qty))

        # 允许某个项目的京能物料全部为客供件；此时仅由ST子件组成ERP BOM。
        return result
    finally:
        wb.close()


# ============================================================
# BOM
# ============================================================

def build_bom_rows(
    groups: List[ParentGroup],
    code_index: Dict[str, Path],
) -> Tuple[List[Tuple[Any, Any, Any, Any]], List[str]]:
    rows: List[Tuple[Any, Any, Any, Any]] = []
    errors: List[str] = []
    cache: Dict[Path, List[Tuple[str, Any]]] = {}

    for group in groups:
        try:
            project_key = resolve_group_project_key(group)

            fillback = code_index.get(project_key.upper())
            if fillback is None:
                raise RuntimeError(
                    f"B列项目号 {project_key} 找不到对应的脚本3输出。"
                )

            if fillback not in cache:
                cache[fillback] = read_jingneng_children(fillback)

            # 同一ET母件下，相同子件统一合并数量。
            # 保持首次出现顺序：ST在前，京能物料在后；若京能物料号与前面的ST相同，
            # 只累加到该ST子件，不新增第二行。
            merged: Dict[str, Tuple[str, Decimal]] = {}

            def add_child(material: str, qty: Any) -> None:
                normalized_material = normalize(material)
                if not normalized_material:
                    return
                key = normalized_material.upper()
                qty_decimal = quantity_to_decimal(qty, normalized_material)
                if key in merged:
                    display_material, old_qty = merged[key]
                    merged[key] = (display_material, old_qty + qty_decimal)
                else:
                    merged[key] = (normalized_material, qty_decimal)

            # ST先加入，数量固定1。
            for st_material in group.st_children:
                add_child(st_material, 1)

            # 再加入京能物料；相同子件在同一母件内自动累加。
            for material, qty in cache[fillback]:
                add_child(material, qty)

            for sequence, (_key, (material, qty_decimal)) in enumerate(merged.items(), start=1):
                rows.append((
                    group.parent_material,
                    material,
                    decimal_to_excel_value(qty_decimal),
                    sequence,
                ))

            source_count = len(group.st_children) + len(cache[fillback])
            merged_count = len(merged)
            log(
                f"[母件] {group.parent_material} -> {project_key} -> "
                f"{fillback.name}"
            )
            log(
                f"       ST {len(group.st_children)} 条，"
                f"京能物料 {len(cache[fillback])} 条，"
                f"合并后子件 {merged_count} 条"
            )
            if merged_count < source_count:
                log(f"       合并重复子件 {source_count - merged_count} 条")

        except Exception as exc:
            errors.append(
                f"ET母件 {group.parent_material}"
                f"（物料导入第 {group.source_row} 行）：{exc}"
            )

    return rows, errors


# ============================================================
# 真正 .xls 输出
# ============================================================

def write_true_xls(
    output_path: Path,
    rows: List[Tuple[Any, Any, Any, Any]],
) -> None:
    if not TEMPLATE_FILE.exists():
        raise RuntimeError(f"找不到ERP xls模板：{TEMPLATE_FILE}")

    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError(
            "缺少 pywin32。请通过启动工具运行。"
        ) from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE_FILE, output_path)

    excel = None
    wb = None

    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False

        wb = excel.Workbooks.Open(str(output_path.resolve()))

        try:
            ws = wb.Worksheets("BomBatch")
        except Exception:
            ws = wb.Worksheets(1)

        last_row = max(int(ws.UsedRange.Rows.Count), 2)
        ws.Range(f"A2:D{last_row}").ClearContents()

        for col, header in enumerate(ERP_HEADERS, start=1):
            ws.Cells(1, col).Value = header

        if rows:
            end_row = len(rows) + 1
            data = tuple(tuple(row) for row in rows)

            ws.Range(f"A2:B{end_row}").NumberFormat = "@"
            ws.Range(f"D2:D{end_row}").NumberFormat = "0"
            ws.Range(f"A2:D{end_row}").Value = data

        wb.Save()

    finally:
        if wb is not None:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass

        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


# ============================================================
# 项目发现
# ============================================================

def project_display_name(project_dir: Path) -> str:
    return project_dir.name


def write_error_report(project_name: str, errors: List[str]) -> Path:
    ERROR_DIR.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[<>:"/\\|?*]+', "_", project_name)
    target = unique_path(ERROR_DIR, f"{safe_name}_ERP导入异常.txt")
    lines = [
        f"项目：{project_name}",
        f"时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
        "本项目未生成 ERP 导入表。",
        "原因：",
        "",
    ]
    lines.extend(f"{i}. {x}" for i, x in enumerate(errors, 1))
    target.write_text("\n".join(lines), encoding="utf-8")
    return target


def process_project(project_dir: Path, code_index: Dict[str, Path]) -> Optional[Path]:
    project_name = project_dir.name
    log("")
    log("=" * 70)
    log(f"项目：{project_name}")
    log("=" * 70)

    try:
        material_import = find_material_import_file(project_dir)
        log(f"物料导入表：{material_import.name}")
        groups = read_material_import(material_import)
        log(f"ET母件：{len(groups)} 个")
    except Exception as exc:
        report = write_error_report(project_name, [str(exc)])
        log(f"[失败] {report}")
        return None

    bom_rows, bom_errors = build_bom_rows(groups, code_index)
    if bom_errors:
        report = write_error_report(project_name, bom_errors)
        log(f"[失败] BOM不完整：{report}")
        return None
    if not bom_rows:
        report = write_error_report(project_name, ["没有生成任何BOM子件记录。"])
        log(f"[失败] {report}")
        return None

    safe_name = re.sub(r'[<>:"/\\|?*]+', "_", project_name)
    output_path = unique_path(OUTPUT_DIR, f"ERP物料清单批量导入表_{safe_name}.xls")
    try:
        write_true_xls(output_path, bom_rows)
    except Exception as exc:
        output_path.unlink(missing_ok=True)
        report = write_error_report(project_name, [str(exc)])
        log(f"[失败] xls生成失败：{report}")
        return None

    log(f"[完成] BOM记录：{len(bom_rows)} 条")
    log(f"[输出] {output_path}")
    return output_path


def run_project(project_dir: Path) -> Dict[str, Any]:
    from project_utils import project_log_dir, project_output_dir
    global OUTPUT_DIR, ERROR_DIR, FILLBACK_OUTPUT_DIR, LOG_FILE

    LOG_LINES.clear()
    project_dir = Path(project_dir).resolve()
    OUTPUT_DIR = project_output_dir(project_dir) / "ERP"
    ERROR_DIR = project_output_dir(project_dir) / "异常"
    FILLBACK_OUTPUT_DIR = project_output_dir(project_dir) / "物料回填"
    LOG_FILE = project_log_dir(project_dir) / "04_生成ERP物料清单.log"
    for d in (OUTPUT_DIR, ERROR_DIR, FILLBACK_OUTPUT_DIR, LOG_FILE.parent):
        d.mkdir(parents=True, exist_ok=True)

    log("=" * 70)
    log("脚本4：生成 ERP 物料清单批量导入表")
    log("=" * 70)
    log(f"项目目录：{project_dir}")
    log(f"物料回填来源：{FILLBACK_OUTPUT_DIR}")
    log(f"ERP xls模板：{TEMPLATE_FILE}")
    log("")

    if not TEMPLATE_FILE.exists():
        result = {"success": False, "can_continue": False, "reason": f"找不到ERP模板：{TEMPLATE_FILE}"}
        log(f"[错误] {result['reason']}")
        save_log()
        return result

    code_index, index_errors = build_fillback_code_index()
    if index_errors:
        reason = "脚本3输出文件名存在项目号歧义：" + "；".join(index_errors)
        log(f"[错误] {reason}")
        save_log()
        return {"success": False, "can_continue": False, "reason": reason, "log_file": str(LOG_FILE)}

    output = process_project(project_dir, code_index)
    save_log()
    if output is None:
        return {"success": False, "can_continue": False, "reason": "ERP生成失败，请查看项目异常报告和日志。", "log_file": str(LOG_FILE)}
    return {"success": True, "can_continue": True, "reason": "", "output_file": str(output), "log_file": str(LOG_FILE)}


def main() -> None:
    import argparse
    from project_utils import add_project_args, ensure_project_dir, write_result_json
    parser = argparse.ArgumentParser(description="脚本4：生成 ERP 物料清单批量导入表")
    add_project_args(parser)
    args = parser.parse_args()
    try:
        result = run_project(ensure_project_dir(args.project))
    except Exception as exc:
        result = {"success": False, "can_continue": False, "reason": str(exc)}
        print(f"[错误] {exc}")
    write_result_json(args.result_json, result)
    raise SystemExit(0 if result.get("success") else 1)


if __name__ == "__main__":
    main()
