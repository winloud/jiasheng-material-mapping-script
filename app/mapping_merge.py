from __future__ import annotations

import shutil
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


# ============================================================
# 配置区
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
MASTER_FILE = DATA_DIR / "嘉盛物料映射汇总表.xlsx"
MERGE_DIR = DATA_DIR / "映射合并"
MERGE_PENDING = MERGE_DIR / "待合并"
MERGE_DONE = MERGE_DIR / "已合并"
BACKUP_DIR = PROJECT_ROOT / "backup"
LOG_DIR = PROJECT_ROOT / "logs"

CONFIG = {
    # 主汇总表
    "master_file": MASTER_FILE,

    # 别人维护过的汇总表放这里
    "merge_dir": MERGE_PENDING,

    # 成功处理后的文件移动到这里
    "done_dir": MERGE_DONE,

    # 主表备份目录
    "backup_dir": BACKUP_DIR,

    # 合并日志
    "log_file": LOG_DIR / "mapping_merge.log",

    # 主匹配列：B = 物料号
    "primary_key_col": "B",

    # B为空时，或B找不到时用于辅助判断：D
    "fallback_key_col": "D",

    # 人工维护的映射区
    "mapping_start_col": "M",
    "mapping_end_col": "V",

    # 覆盖时只允许覆盖人工映射区 M:V
    # 唯一例外：主表B为空、来源表B有值且D唯一匹配时，
    # 可以安全补写主表B列。
    "overwrite_mapping_only": True,

    # 冲突时让用户选择：覆盖 / 新增 / 跳过 / 取消全部
    "interactive_conflict": True,

    # 写入主表前自动完整备份
    "backup_before_merge": True,

    # 若主表映射单元格为空、来源表对应单元格有值：
    # 自动补齐，不提示冲突。
    "auto_fill_blank_mapping_cells": True,

    # 来源表映射单元格为空时，绝不清空主表已有映射。
    "never_clear_mapping_with_blank": True,

    # 支持的待合并文件
    "extensions": {".xlsx", ".xlsm"},
}


# ============================================================
# 日志
# ============================================================

LOG_LINES: List[str] = []


def log(msg: str = "") -> None:
    print(msg)
    LOG_LINES.append(msg)


def save_log() -> None:
    CONFIG["log_file"].write_text(
        "\n".join(LOG_LINES),
        encoding="utf-8",
    )


# ============================================================
# 工具函数
# ============================================================

def is_blank(value: Any) -> bool:
    return value is None or (
        isinstance(value, str) and value.strip() == ""
    )


def normalize_key(value: Any) -> str:
    if is_blank(value):
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def display_value(value: Any) -> str:
    if is_blank(value):
        return "<空>"
    return str(value)


def make_backup(master_path: Path) -> Optional[Path]:
    if not master_path.exists():
        return None

    backup_dir: Path = CONFIG["backup_dir"]
    backup_dir.mkdir(parents=True, exist_ok=True)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    target = backup_dir / f"{master_path.stem}_{stamp}{master_path.suffix}"

    i = 1
    while target.exists():
        target = backup_dir / (
            f"{master_path.stem}_{stamp}_{i}{master_path.suffix}"
        )
        i += 1

    shutil.copy2(master_path, target)
    return target


def unique_destination(folder: Path, source_name: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)

    target = folder / source_name
    if not target.exists():
        return target

    p = Path(source_name)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return folder / f"{p.stem}_{stamp}{p.suffix}"


def select_sheet(workbook):
    if "汇总" in workbook.sheetnames:
        return workbook["汇总"]
    return workbook.active


def validate_header(ws, file_path: Path) -> None:
    primary_col = column_index_from_string(CONFIG["primary_key_col"])
    fallback_col = column_index_from_string(CONFIG["fallback_key_col"])

    b_header = normalize_key(ws.cell(1, primary_col).value)
    d_header = normalize_key(ws.cell(1, fallback_col).value)

    # 这里不要求完全固定，只阻止明显选错表。
    if "料号" not in b_header and "物料" not in b_header:
        raise RuntimeError(
            f"B列表头异常：{display_value(ws.cell(1, primary_col).value)}"
        )

    if not d_header:
        raise RuntimeError("D列表头为空，无法作为辅助匹配字段。")


def build_indexes(ws) -> Tuple[Dict[str, List[int]], Dict[str, List[int]]]:
    primary_col = column_index_from_string(CONFIG["primary_key_col"])
    fallback_col = column_index_from_string(CONFIG["fallback_key_col"])

    by_b: Dict[str, List[int]] = {}
    by_d: Dict[str, List[int]] = {}

    for row_no in range(2, ws.max_row + 1):
        b = normalize_key(ws.cell(row_no, primary_col).value)
        d = normalize_key(ws.cell(row_no, fallback_col).value)

        if b:
            by_b.setdefault(b, []).append(row_no)

        if d:
            by_d.setdefault(d, []).append(row_no)

    return by_b, by_d


def add_index(index: Dict[str, List[int]], key: str, row_no: int) -> None:
    if not key:
        return
    rows = index.setdefault(key, [])
    if row_no not in rows:
        rows.append(row_no)


def remove_index(index: Dict[str, List[int]], key: str, row_no: int) -> None:
    if not key or key not in index:
        return
    rows = index[key]
    if row_no in rows:
        rows.remove(row_no)
    if not rows:
        index.pop(key, None)


def mapping_range() -> Tuple[int, int]:
    return (
        column_index_from_string(CONFIG["mapping_start_col"]),
        column_index_from_string(CONFIG["mapping_end_col"]),
    )


def row_mapping_values(ws, row_no: int) -> List[Any]:
    start_col, end_col = mapping_range()
    return [
        ws.cell(row_no, col).value
        for col in range(start_col, end_col + 1)
    ]


def copy_mapping_safely(
    master_ws,
    master_row: int,
    source_ws,
    source_row: int,
) -> Tuple[int, List[Tuple[int, Any, Any]]]:
    """
    对 M:V 做“安全合并”：
    - 相同：不动
    - 主表空、来源有值：自动补齐
    - 来源空、主表有值：保留主表，不清空
    - 双方都有值且不同：返回冲突，不修改冲突格

    返回：
    (安全自动补齐的单元格数量, 冲突列表)
    """
    start_col, end_col = mapping_range()

    auto_filled = 0
    conflicts: List[Tuple[int, Any, Any]] = []

    for col in range(start_col, end_col + 1):
        old = master_ws.cell(master_row, col).value
        new = source_ws.cell(source_row, col).value

        old_blank = is_blank(old)
        new_blank = is_blank(new)

        if normalize_key(old) == normalize_key(new):
            continue

        if old_blank and not new_blank:
            if CONFIG["auto_fill_blank_mapping_cells"]:
                master_ws.cell(master_row, col).value = new
                auto_filled += 1
            else:
                conflicts.append((col, old, new))
            continue

        if not old_blank and new_blank:
            if CONFIG["never_clear_mapping_with_blank"]:
                continue
            conflicts.append((col, old, new))
            continue

        # 两边都非空且值不同
        conflicts.append((col, old, new))

    return auto_filled, conflicts


def overwrite_mapping(
    master_ws,
    master_row: int,
    source_ws,
    source_row: int,
) -> None:
    """
    用户明确选择“覆盖”时：
    - 只覆盖 M:V；
    - 若来源单元格为空，则默认不清空主表现有值；
    - 不覆盖 A:L 的基础字段。
    """
    start_col, end_col = mapping_range()

    for col in range(start_col, end_col + 1):
        new = source_ws.cell(source_row, col).value

        if CONFIG["never_clear_mapping_with_blank"] and is_blank(new):
            continue

        master_ws.cell(master_row, col).value = new


def append_source_row(
    master_ws,
    source_ws,
    source_row: int,
    next_seq: int,
) -> int:
    """
    新增整行：
    - A列序号使用主表新的连续序号；
    - B:V 复制来源表内容；
    - 超出V列的内容不参与本工具。
    """
    end_col = column_index_from_string(CONFIG["mapping_end_col"])
    new_row = master_ws.max_row + 1

    master_ws.cell(new_row, 1).value = next_seq

    for col in range(2, end_col + 1):
        src = source_ws.cell(source_row, col)
        dst = master_ws.cell(new_row, col)
        dst.value = src.value

        # 复制常见格式，避免新增行看起来完全不同。
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

    return new_row


def max_sequence(ws) -> int:
    result = 0
    for row_no in range(2, ws.max_row + 1):
        value = ws.cell(row_no, 1).value
        try:
            n = int(value)
            result = max(result, n)
        except (TypeError, ValueError):
            pass
    return result


def show_conflict(
    source_file: Path,
    source_ws,
    source_row: int,
    master_ws,
    candidate_rows: List[int],
    reason: str,
) -> None:
    primary_col = column_index_from_string(CONFIG["primary_key_col"])
    fallback_col = column_index_from_string(CONFIG["fallback_key_col"])
    start_col, end_col = mapping_range()

    log("")
    log("=" * 72)
    log("发现异常 / 冲突")
    log(f"原因：{reason}")
    log(f"来源文件：{source_file.name}")
    log(f"来源行：{source_row}")
    log(
        f"来源 B={display_value(source_ws.cell(source_row, primary_col).value)} | "
        f"D={display_value(source_ws.cell(source_row, fallback_col).value)}"
    )

    for row_no in candidate_rows:
        log("-" * 72)
        log(
            f"主表第 {row_no} 行："
            f"B={display_value(master_ws.cell(row_no, primary_col).value)} | "
            f"D={display_value(master_ws.cell(row_no, fallback_col).value)}"
        )

        diffs = []
        for col in range(start_col, end_col + 1):
            old = master_ws.cell(row_no, col).value
            new = source_ws.cell(source_row, col).value
            if normalize_key(old) != normalize_key(new):
                col_letter = master_ws.cell(1, col).column_letter
                header = display_value(master_ws.cell(1, col).value)
                diffs.append(
                    f"{col_letter}({header}): "
                    f"主表={display_value(old)} | "
                    f"来源={display_value(new)}"
                )

        if diffs:
            for diff in diffs:
                log("  " + diff)
        else:
            log("  M:V 映射区没有差异；冲突来自匹配字段。")

    log("=" * 72)


def choose_candidate_row(
    candidate_rows: List[int],
) -> Optional[int]:
    if len(candidate_rows) == 1:
        return candidate_rows[0]

    while True:
        raw = input(
            f"请选择要覆盖的主表行号 {candidate_rows}，"
            "或输入 S 放弃覆盖："
        ).strip()

        if raw.lower() == "s":
            return None

        try:
            row_no = int(raw)
        except ValueError:
            print("输入无效。")
            continue

        if row_no in candidate_rows:
            return row_no

        print("该行号不在候选列表中。")


def ask_action(
    apply_all_action: Optional[str],
) -> Tuple[str, Optional[str]]:
    if apply_all_action:
        return apply_all_action, apply_all_action

    while True:
        print(
            "\n请选择："
            "\n  [O]  覆盖映射（仅M:V；不会用空值清空主表）"
            "\n  [A]  新增为一行"
            "\n  [S]  跳过"
            "\n  [Q]  取消本次全部合并，不保存主表"
            "\n  [OA] 后续冲突全部覆盖"
            "\n  [AA] 后续冲突全部新增"
            "\n  [SA] 后续冲突全部跳过"
        )
        choice = input("输入选择：").strip().upper()

        if choice in {"O", "A", "S", "Q"}:
            return choice, None

        if choice == "OA":
            return "O", "O"

        if choice == "AA":
            return "A", "A"

        if choice == "SA":
            return "S", "S"

        print("输入无效，请重新选择。")


# ============================================================
# 主程序
# ============================================================

def main() -> None:
    master_path = Path(CONFIG["master_file"])
    merge_dir: Path = CONFIG["merge_dir"]
    done_dir: Path = CONFIG["done_dir"]

    merge_dir.mkdir(parents=True, exist_ok=True)
    done_dir.mkdir(parents=True, exist_ok=True)
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    LOG_DIR.mkdir(parents=True, exist_ok=True)

    LOG_LINES.clear()

    log("=" * 72)
    log("合并别人维护过的映射汇总表")
    log("=" * 72)
    log(f"主表：{master_path}")
    log(f"待合并目录：{merge_dir}")

    if not master_path.exists():
        log(f"[错误] 找不到主汇总表：{master_path}")
        save_log()
        return

    source_files = sorted(
        p for p in merge_dir.iterdir()
        if p.is_file()
        and p.suffix.lower() in CONFIG["extensions"]
        and p.name != master_path.name
        and not p.name.startswith("~$")
    )

    if not source_files:
        log("[提示] 待合并目录中没有可处理的 xlsx/xlsm 文件。")
        save_log()
        return

    try:
        master_wb = load_workbook(master_path)
    except PermissionError:
        log("[错误] 主汇总表正在被 Excel 占用，请关闭后重试。")
        save_log()
        return

    master_ws = select_sheet(master_wb)

    try:
        validate_header(master_ws, master_path)
    except Exception as exc:
        log(f"[错误] 主汇总表结构检查失败：{exc}")
        master_wb.close()
        save_log()
        return

    by_b, by_d = build_indexes(master_ws)
    next_seq = max_sequence(master_ws)

    stats = {
        "files_ok": 0,
        "files_error": 0,
        "rows_scanned": 0,
        "rows_added": 0,
        "rows_auto_merged": 0,
        "rows_overwritten": 0,
        "rows_skipped": 0,
        "conflicts": 0,
        "b_backfilled": 0,
    }

    processed_files: List[Path] = []
    apply_all_action: Optional[str] = None
    cancelled = False
    any_change = False

    primary_col = column_index_from_string(CONFIG["primary_key_col"])
    fallback_col = column_index_from_string(CONFIG["fallback_key_col"])

    for source_file in source_files:
        log("")
        log(f"[文件] {source_file.name}")

        try:
            source_wb = load_workbook(
                source_file,
                data_only=False,
                read_only=False,
            )
            source_ws = select_sheet(source_wb)
            validate_header(source_ws, source_file)
        except Exception as exc:
            stats["files_error"] += 1
            log(f"  [错误] 无法读取或表结构异常：{exc}")
            try:
                source_wb.close()
            except Exception:
                pass
            continue

        file_cancelled = False

        for source_row in range(2, source_ws.max_row + 1):
            b = normalize_key(source_ws.cell(source_row, primary_col).value)
            d = normalize_key(source_ws.cell(source_row, fallback_col).value)

            if not b and not d:
                continue

            stats["rows_scanned"] += 1

            candidate_rows: List[int] = []
            reason = ""

            # ------------------------------------------------
            # 1. B有值：优先按B匹配
            # ------------------------------------------------
            if b:
                b_rows = list(by_b.get(b, []))

                if len(b_rows) == 1:
                    candidate_rows = b_rows
                    master_row = b_rows[0]
                    master_d = normalize_key(
                        master_ws.cell(master_row, fallback_col).value
                    )

                    # 同一B，但双方D都有值且不同：异常
                    if d and master_d and d != master_d:
                        reason = "同一B物料号，但D列内容不同"
                    else:
                        auto_filled, mapping_conflicts = copy_mapping_safely(
                            master_ws,
                            master_row,
                            source_ws,
                            source_row,
                        )

                        if mapping_conflicts:
                            reason = "同一物料号的人工映射 M:V 存在冲突"
                        else:
                            if auto_filled:
                                any_change = True
                                stats["rows_auto_merged"] += 1
                                log(
                                    f"  [自动合并] 来源第{source_row}行 -> "
                                    f"主表第{master_row}行，"
                                    f"补齐映射单元格 {auto_filled} 个"
                                )
                            else:
                                stats["rows_skipped"] += 1
                            continue

                elif len(b_rows) > 1:
                    candidate_rows = b_rows
                    reason = "主表中同一B物料号存在多行"

                else:
                    # B不存在，再按D判断。
                    d_rows = list(by_d.get(d, [])) if d else []

                    if not d_rows:
                        next_seq += 1
                        new_row = append_source_row(
                            master_ws,
                            source_ws,
                            source_row,
                            next_seq,
                        )
                        add_index(by_b, b, new_row)
                        add_index(by_d, d, new_row)
                        any_change = True
                        stats["rows_added"] += 1
                        log(
                            f"  [新增] 来源第{source_row}行 -> "
                            f"主表第{new_row}行，B={b}"
                        )
                        continue

                    # 若D唯一匹配到主表B为空行，可安全补B并合并映射。
                    blank_b_rows = [
                        r for r in d_rows
                        if not normalize_key(
                            master_ws.cell(r, primary_col).value
                        )
                    ]

                    nonblank_other_b_rows = [
                        r for r in d_rows
                        if normalize_key(
                            master_ws.cell(r, primary_col).value
                        )
                        and normalize_key(
                            master_ws.cell(r, primary_col).value
                        ) != b
                    ]

                    if (
                        len(d_rows) == 1
                        and len(blank_b_rows) == 1
                        and not nonblank_other_b_rows
                    ):
                        master_row = blank_b_rows[0]
                        master_ws.cell(master_row, primary_col).value = (
                            source_ws.cell(source_row, primary_col).value
                        )
                        add_index(by_b, b, master_row)
                        stats["b_backfilled"] += 1
                        any_change = True

                        auto_filled, mapping_conflicts = copy_mapping_safely(
                            master_ws,
                            master_row,
                            source_ws,
                            source_row,
                        )

                        if mapping_conflicts:
                            candidate_rows = [master_row]
                            reason = (
                                "按D唯一匹配到主表B为空行，已安全补B；"
                                "但人工映射 M:V 存在冲突"
                            )
                        else:
                            if auto_filled:
                                stats["rows_auto_merged"] += 1
                            log(
                                f"  [补B并合并] 来源第{source_row}行 -> "
                                f"主表第{master_row}行，B={b}"
                            )
                            continue

                    else:
                        candidate_rows = d_rows
                        if nonblank_other_b_rows:
                            reason = "B在主表不存在，但同一D对应了其它非空B物料号"
                        else:
                            reason = "B在主表不存在，但D在主表存在多重匹配"

            # ------------------------------------------------
            # 2. B为空：按D匹配
            # ------------------------------------------------
            else:
                d_rows = list(by_d.get(d, []))

                if not d_rows:
                    next_seq += 1
                    new_row = append_source_row(
                        master_ws,
                        source_ws,
                        source_row,
                        next_seq,
                    )
                    add_index(by_d, d, new_row)
                    any_change = True
                    stats["rows_added"] += 1
                    log(
                        f"  [新增] 来源第{source_row}行 -> "
                        f"主表第{new_row}行，B为空，D={d}"
                    )
                    continue

                if len(d_rows) == 1:
                    candidate_rows = d_rows
                    master_row = d_rows[0]

                    auto_filled, mapping_conflicts = copy_mapping_safely(
                        master_ws,
                        master_row,
                        source_ws,
                        source_row,
                    )

                    if mapping_conflicts:
                        reason = "B为空，按D匹配后人工映射 M:V 存在冲突"
                    else:
                        if auto_filled:
                            any_change = True
                            stats["rows_auto_merged"] += 1
                            log(
                                f"  [自动合并] 来源第{source_row}行 -> "
                                f"主表第{master_row}行，"
                                f"补齐映射单元格 {auto_filled} 个"
                            )
                        else:
                            stats["rows_skipped"] += 1
                        continue
                else:
                    candidate_rows = d_rows
                    reason = "B为空，D在主表中存在多重匹配"

            # ------------------------------------------------
            # 3. 异常 / 冲突：用户决策
            # ------------------------------------------------
            stats["conflicts"] += 1
            show_conflict(
                source_file,
                source_ws,
                source_row,
                master_ws,
                candidate_rows,
                reason,
            )

            if not CONFIG["interactive_conflict"]:
                action = "S"
                new_apply_all = None
            else:
                action, new_apply_all = ask_action(apply_all_action)
                if new_apply_all is not None:
                    apply_all_action = new_apply_all

            if action == "Q":
                cancelled = True
                file_cancelled = True
                break

            if action == "S":
                stats["rows_skipped"] += 1
                log("  [决定] 跳过")
                continue

            if action == "A":
                next_seq += 1
                new_row = append_source_row(
                    master_ws,
                    source_ws,
                    source_row,
                    next_seq,
                )
                add_index(by_b, b, new_row)
                add_index(by_d, d, new_row)
                any_change = True
                stats["rows_added"] += 1
                log(f"  [决定] 新增 -> 主表第{new_row}行")
                continue

            if action == "O":
                target_row = choose_candidate_row(candidate_rows)
                if target_row is None:
                    stats["rows_skipped"] += 1
                    log("  [决定] 放弃覆盖，本条跳过")
                    continue

                # 覆盖只改 M:V。
                overwrite_mapping(
                    master_ws,
                    target_row,
                    source_ws,
                    source_row,
                )

                # 唯一允许碰基础区的情况：
                # 主表B为空 + 来源B有值，则补写B，不覆盖已有非空B。
                old_b = normalize_key(
                    master_ws.cell(target_row, primary_col).value
                )
                if not old_b and b:
                    master_ws.cell(target_row, primary_col).value = (
                        source_ws.cell(source_row, primary_col).value
                    )
                    add_index(by_b, b, target_row)
                    stats["b_backfilled"] += 1

                any_change = True
                stats["rows_overwritten"] += 1
                log(
                    f"  [决定] 覆盖映射 -> 主表第{target_row}行 "
                    f"(仅M:V，基础字段不覆盖)"
                )

        source_wb.close()

        if file_cancelled:
            break

        stats["files_ok"] += 1
        processed_files.append(source_file)

    # --------------------------------------------------------
    # Q：不保存主表，不移动任何来源文件
    # --------------------------------------------------------
    if cancelled:
        master_wb.close()
        log("")
        log("=" * 72)
        log("用户取消：本次所有主表修改均未保存，来源文件也未移动。")
        log("=" * 72)
        save_log()
        return

    # --------------------------------------------------------
    # 保存：先备份，再一次性写主表
    # --------------------------------------------------------
    if any_change:
        try:
            if CONFIG["backup_before_merge"]:
                backup_path = make_backup(master_path)
                if backup_path:
                    log(f"[备份] {backup_path}")

            master_wb.save(master_path)
            log(f"[保存] 主汇总表已更新：{master_path}")

        except PermissionError:
            master_wb.close()
            log(
                "[错误] 保存失败：主汇总表可能正在被 Excel 占用。"
                "来源文件未移动。"
            )
            save_log()
            return
        except Exception as exc:
            master_wb.close()
            log(f"[错误] 保存主汇总表失败：{exc}。来源文件未移动。")
            save_log()
            return
    else:
        master_wb.close()
        log("[提示] 没有需要写入主汇总表的变化。")

    # --------------------------------------------------------
    # 成功处理过的来源文件移动到“已合并”
    # --------------------------------------------------------
    moved = 0
    for source_file in processed_files:
        try:
            target = unique_destination(done_dir, source_file.name)
            shutil.move(str(source_file), str(target))
            moved += 1
            log(f"[已合并] {source_file.name} -> {target.name}")
        except Exception as exc:
            log(f"[警告] 无法移动 {source_file.name}：{exc}")

    # --------------------------------------------------------
    # 汇总
    # --------------------------------------------------------
    log("")
    log("=" * 72)
    log("合并完成")
    log("=" * 72)
    log(f"成功读取文件：{stats['files_ok']}")
    log(f"读取异常文件：{stats['files_error']}")
    log(f"扫描有效行：{stats['rows_scanned']}")
    log(f"新增行：{stats['rows_added']}")
    log(f"自动补齐映射行：{stats['rows_auto_merged']}")
    log(f"人工选择覆盖：{stats['rows_overwritten']}")
    log(f"安全补写B物料号：{stats['b_backfilled']}")
    log(f"冲突/异常：{stats['conflicts']}")
    log(f"跳过行：{stats['rows_skipped']}")
    log(f"移动到已合并：{moved}")
    log(f"日志：{CONFIG['log_file']}")

    save_log()


if __name__ == "__main__":
    main()
