from __future__ import annotations

import copy
import re
import shutil
import tempfile
import time
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment


# ============================================================
# 项目路径
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
MASTER_FILE = DATA_DIR / "嘉盛物料映射汇总表.xlsx"

FILLBACK_DIR = DATA_DIR / "物料回填"
INPUT_DIR = FILLBACK_DIR / "待处理"
OUTPUT_DIR = FILLBACK_DIR / "输出"
DONE_DIR = FILLBACK_DIR / "已处理"
ERROR_DIR = FILLBACK_DIR / "异常"
LOG_FILE = PROJECT_ROOT / "logs" / "material_fillback.log"


# ============================================================
# 配置
# ============================================================

CONFIG = {
    # 主映射表中客户物料匹配字段
    "master_material_col": "B",
    "master_model_col": "D",

    # 映射区：物料A~物料F，共 6 组，每组“我方物料号 + 单位映射数量”
    # M:N, O:P, Q:R, S:T, U:V, W:X
    "mapping_start_col": "M",
    "mapping_count": 6,

    # 回填目标列
    "result_material_col": "AA",
    "result_qty_col": "AB",
    "result_status_col": "AC",

    # 客户清单 Sheet 规则
    "sheets": [
        {
            "name": "电柜元器件",
            "start_row": 4,
            "stop_col": "D",
            "material_col": "B",
            "model_col": "D",
            "quantity_col": "F",
        },
        {
            "name": "电柜半成品",
            "start_row": 3,
            "stop_col": "D",
            "material_col": "B",
            "model_col": "D",
            "quantity_col": "G",
        },
    ],

    # 异常行高亮范围
    "highlight_start_col": "A",
    "highlight_end_col": "AC",

    # 后续脚本统一读取的标准化汇总 Sheet
    "summary_sheet_name": "京能物料对应汇总",

    # 主映射表中的公共物料 Sheet
    "common_material_sheet_name": "公共物料",
}


# ============================================================
# 报价单驱动公共物料数量
# ============================================================

# 柜号示例：JSA26125-11~13、HZJF26071-26~29-47~50-55
NUMBER_TOKEN_RE = re.compile(r"^(?P<start>\d+)(?:~(?P<end>\d+))?(?:\([^)]*\))?$")
BASE_RE = re.compile(r"^[A-Za-z]+\d+$")

# 特殊公共物料规则直接维护在 Python 中。
# key = 公共物料 Sheet D 列型号；quote_model = 报价明细 C 列实际型号。
QUOTE_COMMON_MATERIAL_RULES = {
    "ZQV2.5N/50": {
        "quote_model": "ZQV2.5N/10",
        "quantity_divisor": Decimal("5"),
        "round_after_division": True,
    },
}


# ============================================================
# 日志
# ============================================================

_log_lines: List[str] = []


def log(message: str = "") -> None:
    print(message)
    _log_lines.append(message)


def save_log() -> None:
    LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    LOG_FILE.write_text("\n".join(_log_lines), encoding="utf-8")


# ============================================================
# 通用工具
# ============================================================


def is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def normalize_key(value: Any) -> str:
    """标准化 B/D 匹配键，避免 123 / 123.0 / '123' 类型差异。"""
    if is_empty(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_decimal(value: Any) -> Optional[Decimal]:
    """空值返回 None；合法数字返回 Decimal；非法值抛 ValueError。"""
    if is_empty(value):
        return None
    if isinstance(value, bool):
        raise ValueError(f"布尔值不是合法数量：{value}")
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError(f"不是合法数字：{value}")


def decimal_to_excel(value: Decimal) -> Any:
    """整数以 int 写回 Excel，小数以 float 写回。"""
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def col_to_num(col: str) -> int:
    result = 0
    for ch in col.strip().upper():
        result = result * 26 + ord(ch) - ord("A") + 1
    return result


# 高频列号预计算，避免在逐行处理循环中重复解析 Excel 列名。
RESULT_MATERIAL_COL = col_to_num(CONFIG["result_material_col"])
RESULT_QTY_COL = col_to_num(CONFIG["result_qty_col"])
RESULT_STATUS_COL = col_to_num(CONFIG["result_status_col"])
HIGHLIGHT_START_COL = col_to_num(CONFIG["highlight_start_col"])
HIGHLIGHT_END_COL = col_to_num(CONFIG["highlight_end_col"])
MAPPING_START_COL = col_to_num(CONFIG["mapping_start_col"])
MAPPING_MAX_COL = col_to_num("X")


def unique_path(directory: Path, filename: str) -> Path:
    """避免覆盖已有输出/归档文件。"""
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / filename
    if not target.exists():
        return target

    p = Path(filename)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return directory / f"{p.stem}_{stamp}{p.suffix}"


def copy_cell(src, dst) -> None:
    """复制单元格内容和格式。_style 已包含字体/填充/边框/对齐等，避免重复 deepcopy。"""
    dst.value = src.value
    if src.has_style:
        dst._style = copy.copy(src._style)
    if src.hyperlink:
        dst._hyperlink = copy.copy(src.hyperlink)
    if src.comment:
        dst.comment = copy.copy(src.comment)


def copy_row_a_to_z(ws, source_row: int, target_row: int) -> None:
    """1:N 展开时，将原客户行 A:Z 完整复制到新行。"""
    for col in range(1, 27):
        copy_cell(ws.cell(source_row, col), ws.cell(target_row, col))

    source_dim = ws.row_dimensions[source_row]
    target_dim = ws.row_dimensions[target_row]
    if source_dim.height is not None:
        target_dim.height = source_dim.height
    target_dim.hidden = source_dim.hidden
    target_dim.outlineLevel = source_dim.outlineLevel


def material_no_as_text(value: Any) -> Any:
    """
    物料号写入 Excel 时按文本处理，避免长纯数字显示为科学计数法。
    纯数字物料号等效于在 Excel 中手工输入前置单引号。
    """
    if is_empty(value):
        return None

    text = normalize_key(value)

    # 纯数字物料号强制作为文本保存。
    if text.isdigit():
        return text

    return value


def write_material_no(cell, value: Any) -> None:
    """将物料号写成文本格式，避免 Excel 自动转科学计数法。"""
    cell.value = material_no_as_text(value)
    cell.number_format = "@"


def clear_result_cells(ws, row: int) -> None:
    for col in (RESULT_MATERIAL_COL, RESULT_QTY_COL, RESULT_STATUS_COL):
        ws.cell(row, col).value = None


def set_status(ws, row: int, status: str) -> None:
    ws.cell(row, RESULT_STATUS_COL).value = status


def highlight_error_row(ws, row: int) -> None:
    """异常行整行 A:AC 淡红高亮。"""
    fill = PatternFill(fill_type="solid", fgColor="F4CCCC")
    font = Font(color="9C0006")
    for col in range(HIGHLIGHT_START_COL, HIGHLIGHT_END_COL + 1):
        cell = ws.cell(row, col)
        cell.fill = copy.copy(fill)
        # 保留原字体的其他属性，仅对 AC 状态列用红字，避免破坏客户表字体样式。
        if col == RESULT_STATUS_COL:
            cell.font = copy.copy(font)


def apply_result_headers(ws, header_row: int, style_source_col: int) -> None:
    headers = {
        RESULT_MATERIAL_COL: "京能物料号",
        RESULT_QTY_COL: "京能数量",
        RESULT_STATUS_COL: "匹配状态",
    }

    source_cell = ws.cell(header_row, style_source_col)
    for col, text in headers.items():
        target = ws.cell(header_row, col)
        target.value = text
        if source_cell.has_style:
            target._style = copy.copy(source_cell._style)

    # 给新增列一个可读的默认宽度，但不改客户原列宽。
    ws.column_dimensions[CONFIG["result_material_col"]].width = max(
        ws.column_dimensions[CONFIG["result_material_col"]].width or 0, 18
    )
    ws.column_dimensions[CONFIG["result_qty_col"]].width = max(
        ws.column_dimensions[CONFIG["result_qty_col"]].width or 0, 12
    )
    ws.column_dimensions[CONFIG["result_status_col"]].width = max(
        ws.column_dimensions[CONFIG["result_status_col"]].width or 0, 22
    )


# ============================================================
# 标准化汇总 Sheet
# ============================================================


def build_summary_record(
    ws,
    sheet_cfg: Dict[str, Any],
    original_row: int,
    our_material: Any,
    our_qty: Any,
    status: str,
) -> Dict[str, Any]:
    """将两个客户 Sheet 的不同结构统一成标准记录。"""
    sheet_name = sheet_cfg["name"]

    customer_material = ws.cell(original_row, col_to_num(sheet_cfg["material_col"])).value
    customer_name = ws.cell(original_row, 3).value
    customer_model = ws.cell(original_row, col_to_num(sheet_cfg["model_col"])).value
    customer_brand = ws.cell(original_row, 5).value
    customer_qty = ws.cell(original_row, col_to_num(sheet_cfg["quantity_col"])).value

    # 电柜元器件 F 列本身就是数量，没有独立规格列；
    # 电柜半成品 F 列为规格、G 列为数量。
    customer_spec = None
    if sheet_name == "电柜半成品":
        customer_spec = ws.cell(original_row, 6).value

    return {
        "来源Sheet": sheet_name,
        "原始行号": original_row,
        "客户物料号": customer_material,
        "品名": customer_name,
        "型号": customer_model,
        "品牌": customer_brand,
        "规格": customer_spec,
        "客户数量": customer_qty,
        "京能物料号": our_material,
        "京能数量": our_qty,
        "匹配状态": status,
        "物料来源": "客户映射",
    }


def extract_mapping_pairs_from_values(
    values: tuple,
) -> Tuple[List[Tuple[Any, Decimal]], Optional[str]]:
    """
    从 iter_rows(values_only=True) 得到的一整行中读取 M:X 的物料A~F。
    用于主映射表流式扫描，避免 read_only 工作表反复随机访问单元格。
    """
    start_index = MAPPING_START_COL - 1
    pairs: List[Tuple[Any, Decimal]] = []

    for i in range(CONFIG["mapping_count"]):
        material_index = start_index + i * 2
        qty_index = material_index + 1

        material = values[material_index] if material_index < len(values) else None
        mapping_qty = values[qty_index] if qty_index < len(values) else None

        if is_empty(material):
            continue

        if is_empty(mapping_qty):
            factor = Decimal("1")
        else:
            try:
                factor = to_decimal(mapping_qty)
            except ValueError:
                label = chr(ord("A") + i)
                return [], f"物料{label}数量异常：{mapping_qty}"
            assert factor is not None

        pairs.append((material, factor))

    return pairs, None


def build_common_material_definitions(master_ws) -> List[Dict[str, Any]]:
    """读取“公共物料”配置。数量不在这里确定，脚本3按每张电气清单对应的报价柜型块动态取得。"""
    definitions: List[Dict[str, Any]] = []
    max_col = MAPPING_MAX_COL

    for row_no, values in enumerate(
        master_ws.iter_rows(min_row=2, max_col=max_col, values_only=True),
        start=2,
    ):
        name = values[2] if len(values) > 2 else None       # C
        model = values[3] if len(values) > 3 else None      # D
        brand = values[4] if len(values) > 4 else None      # E
        spec = values[5] if len(values) > 5 else None       # F
        pairs, mapping_error = extract_mapping_pairs_from_values(values)

        if is_empty(name) and is_empty(model) and not pairs and not mapping_error:
            continue

        definitions.append({
            "来源Sheet": CONFIG["common_material_sheet_name"],
            "原始行号": row_no,
            "品名": name,
            "型号": model,
            "品牌": brand,
            "规格": spec,
            "pairs": pairs,
            "mapping_error": mapping_error,
        })

    return definitions


def normalize_model(value: Any) -> str:
    """报价型号匹配：忽略首尾空白，并将连续空白/换行折叠为一个空格。"""
    if is_empty(value):
        return ""
    return " ".join(str(value).strip().split()).upper()


def parse_cabinet_expression(value: Any) -> Tuple[str, List[str]]:
    """
    将“JSA26125-11~13”或电气清单文件名前缀展开为单柜集合。
    遇到第一个非编号 token 即停止，因此后续“-电气清单...”不会干扰。
    """
    text = str(value or "").strip()
    if not text:
        return "", []
    # 对文件名输入去掉扩展名；柜号单元格通常无扩展名，此操作同样安全。
    stem = Path(text).stem.strip()
    parts = stem.split("-")
    if len(parts) < 2:
        return "", []

    base = parts[0].strip()
    if not BASE_RE.fullmatch(base):
        return "", []

    codes: List[str] = []
    for token in parts[1:]:
        token = token.strip()
        match = NUMBER_TOKEN_RE.fullmatch(token)
        if not match:
            break
        start_text = match.group("start")
        end_text = match.group("end")
        if end_text is None:
            codes.append(f"{base}-{start_text}")
            continue
        start_num = int(start_text)
        end_num = int(end_text)
        if end_num < start_num:
            raise ValueError(f"反向柜号范围：{token}")
        width = max(len(start_text), len(end_text))
        for number in range(start_num, end_num + 1):
            codes.append(f"{base}-{number:0{width}d}")
    return base, codes


def find_quote_file(project_dir: Path, project_base: str) -> Path:
    """报价单位于项目根目录，文件名必须同时包含项目号和“报价”。"""
    matches = []
    base_upper = project_base.upper()
    for path in project_dir.iterdir():
        if not path.is_file() or path.name.startswith("~$"):
            continue
        if path.suffix.lower() not in {".xls", ".xlsx", ".xlsm"}:
            continue
        name_upper = path.name.upper()
        if base_upper in name_upper and "报价" in path.name:
            matches.append(path)

    if len(matches) == 0:
        raise RuntimeError(f"未找到报价单：文件名需同时包含 {project_base} 和“报价”。")
    if len(matches) > 1:
        raise RuntimeError(
            f"找到多个 {project_base} 报价单，无法自动判断："
            + " / ".join(p.name for p in matches)
        )
    return matches[0]


def load_quote_blocks(path: Path, project_base: str) -> List[Dict[str, Any]]:
    """读取报价单指定项目 Sheet，并解析所有“柜号:”明细块。"""
    suffix = path.suffix.lower()
    rows: List[Tuple[int, Any, Any, Any, Any]] = []  # row, A, B, C, F

    if suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            # 优先精确 Sheet 名；再允许大小写唯一匹配。
            actual_sheet = project_base if project_base in wb.sheetnames else None
            if actual_sheet is None:
                candidates = [n for n in wb.sheetnames if n.upper() == project_base.upper()]
                if len(candidates) == 1:
                    actual_sheet = candidates[0]
            if actual_sheet is None:
                raise RuntimeError(f"报价单中找不到 Sheet：{project_base}")
            ws = wb[actual_sheet]
            for row_no, values in enumerate(
                ws.iter_rows(min_row=1, max_col=6, values_only=True), start=1
            ):
                vals = list(values) + [None] * (6 - len(values))
                rows.append((row_no, vals[0], vals[1], vals[2], vals[5]))
        finally:
            wb.close()
    elif suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("读取 .xls 报价单需要 xlrd。") from exc
        wb = xlrd.open_workbook(str(path), on_demand=True)
        try:
            names = wb.sheet_names()
            actual_sheet = project_base if project_base in names else None
            if actual_sheet is None:
                candidates = [n for n in names if n.upper() == project_base.upper()]
                if len(candidates) == 1:
                    actual_sheet = candidates[0]
            if actual_sheet is None:
                raise RuntimeError(f"报价单中找不到 Sheet：{project_base}")
            ws = wb.sheet_by_name(actual_sheet)
            for idx in range(ws.nrows):
                def cv(col: int):
                    return ws.cell_value(idx, col) if col < ws.ncols else None
                rows.append((idx + 1, cv(0), cv(1), cv(2), cv(5)))
        finally:
            wb.release_resources()
    else:
        raise RuntimeError(f"不支持的报价单格式：{path.suffix}")

    starts: List[Tuple[int, int, str, List[str]]] = []  # list_index, excel_row, raw_group, codes
    for idx, (row_no, a, b, _c, _f) in enumerate(rows):
        a_text = str(a or "").strip().replace("：", ":")
        if not a_text.startswith("柜号"):
            continue
        raw_group = str(b or "").strip()
        base, codes = parse_cabinet_expression(raw_group)
        if base.upper() != project_base.upper() or not codes:
            continue
        starts.append((idx, row_no, raw_group, codes))

    if not starts:
        raise RuntimeError(f"报价单 Sheet {project_base} 中没有找到有效“柜号:”明细块。")

    blocks: List[Dict[str, Any]] = []
    for pos, (start_idx, start_row, raw_group, codes) in enumerate(starts):
        end_idx = starts[pos + 1][0] if pos + 1 < len(starts) else len(rows)
        model_rows: Dict[str, List[Tuple[int, Any, Any]]] = defaultdict(list)
        for row_no, _a, _b, c, f in rows[start_idx + 1:end_idx]:
            model_key = normalize_model(c)
            if not model_key:
                continue
            model_rows[model_key].append((row_no, c, f))
        blocks.append({
            "start_row": start_row,
            "group_text": raw_group,
            "codes": [c.upper() for c in codes],
            "models": dict(model_rows),
        })
    return blocks


def resolve_quote_block(
    source: Path,
    project_dir: Path,
    quote_cache: Dict[Tuple[str, str], Tuple[Path, List[Dict[str, Any]]]],
) -> Tuple[Path, Dict[str, Any]]:
    """按电气清单合并柜号集合，精确匹配报价单中的同一柜型块。"""
    project_base, source_codes = parse_cabinet_expression(source.stem)
    if not project_base or not source_codes:
        raise RuntimeError(f"无法从电气清单文件名解析柜号范围：{source.name}")

    cache_key = (str(project_dir.resolve()), project_base.upper())
    cached = quote_cache.get(cache_key)
    if cached is None:
        quote_file = find_quote_file(project_dir, project_base)
        blocks = load_quote_blocks(quote_file, project_base)
        quote_cache[cache_key] = (quote_file, blocks)
    else:
        quote_file, blocks = cached

    wanted = {c.upper() for c in source_codes}
    matches = [b for b in blocks if set(b["codes"]) == wanted]
    if len(matches) == 0:
        raise RuntimeError(
            f"报价单 {quote_file.name} / Sheet {project_base} 中找不到与电气清单完全对应的柜号块："
            f"{', '.join(source_codes)}"
        )
    if len(matches) > 1:
        raise RuntimeError(
            f"报价单中柜号集合重复出现 {len(matches)} 次，无法唯一定位：{', '.join(source_codes)}"
        )
    return quote_file, matches[0]


def apply_quote_quantity_rule(common_model: Any, quote_qty: Decimal) -> Decimal:
    """公共物料特殊换算。明确使用 Python 内置 round()，不是 roundup。"""
    key = normalize_model(common_model)
    rule = QUOTE_COMMON_MATERIAL_RULES.get(key)
    if not rule:
        return quote_qty

    divisor = rule.get("quantity_divisor")
    if divisor:
        quote_qty = quote_qty / divisor
    if rule.get("round_after_division"):
        quote_qty = Decimal(round(quote_qty))
    return quote_qty


def quote_lookup_model(common_model: Any) -> str:
    key = normalize_model(common_model)
    rule = QUOTE_COMMON_MATERIAL_RULES.get(key)
    if rule:
        return str(rule.get("quote_model") or common_model)
    return str(common_model or "")


def build_common_records_from_quote(
    definitions: List[Dict[str, Any]],
    quote_block: Dict[str, Any],
) -> Tuple[List[Dict[str, Any]], int]:
    """按当前“合并柜型”报价块，为公共物料计算单台柜 BOM 数量。"""
    records: List[Dict[str, Any]] = []
    abnormal = 0
    models = quote_block["models"]

    for definition in definitions:
        base = {
            "来源Sheet": CONFIG["common_material_sheet_name"],
            "原始行号": definition["原始行号"],
            "客户物料号": None,
            "品名": definition.get("品名"),
            "型号": definition.get("型号"),
            "品牌": definition.get("品牌"),
            "规格": definition.get("规格"),
            "客户数量": None,
            "物料来源": "公共物料",
        }
        mapping_error = definition.get("mapping_error")
        pairs = definition.get("pairs") or []

        if mapping_error:
            records.append({**base, "京能物料号": None, "京能数量": None,
                            "匹配状态": mapping_error})
            abnormal += 1
            continue
        if not pairs:
            records.append({**base, "京能物料号": None, "京能数量": None,
                            "匹配状态": "未配置京能物料"})
            abnormal += 1
            continue

        common_model = definition.get("型号")
        lookup_model = quote_lookup_model(common_model)
        lookup_key = normalize_model(lookup_model)
        if not lookup_key:
            records.append({**base, "京能物料号": None, "京能数量": None,
                            "匹配状态": "公共物料型号为空，无法匹配报价"})
            abnormal += 1
            continue

        hits = models.get(lookup_key, [])
        if len(hits) == 0:
            records.append({**base, "京能物料号": None, "京能数量": None,
                            "匹配状态": f"报价未找到型号：{lookup_model}"})
            abnormal += 1
            continue
        if len(hits) > 1:
            rows = ",".join(str(hit[0]) for hit in hits)
            records.append({**base, "京能物料号": None, "京能数量": None,
                            "匹配状态": f"报价型号重复：{lookup_model}（行{rows}）"})
            abnormal += 1
            continue

        quote_row, _raw_model, raw_qty = hits[0]
        try:
            quote_qty = to_decimal(raw_qty)
        except ValueError:
            quote_qty = None
        if quote_qty is None:
            records.append({**base, "京能物料号": None, "京能数量": None,
                            "匹配状态": f"报价数量异常：{lookup_model} / F{quote_row}={raw_qty}"})
            abnormal += 1
            continue

        try:
            effective_qty = apply_quote_quantity_rule(common_model, quote_qty)
        except Exception as exc:
            records.append({**base, "京能物料号": None, "京能数量": None,
                            "匹配状态": f"报价数量换算失败：{lookup_model} / {exc}"})
            abnormal += 1
            continue

        for material, factor in pairs:
            final_qty = effective_qty * factor
            records.append({
                **base,
                "京能物料号": material,
                "京能数量": decimal_to_excel(final_qty),
                "匹配状态": "已匹配",
            })

    return records, abnormal

def create_summary_sheet(wb, records: List[Dict[str, Any]]) -> None:
    """重建“京能物料对应汇总”，作为第4个及后续脚本的统一数据接口。"""
    sheet_name = CONFIG["summary_sheet_name"]
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

    ws = wb.create_sheet(sheet_name)
    headers = [
        "来源Sheet", "原始行号", "客户物料号", "品名", "型号", "品牌",
        "规格", "客户数量", "京能物料号", "京能数量", "匹配状态", "物料来源",
    ]
    ws.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    error_fill = PatternFill(fill_type="solid", fgColor="F4CCCC")
    error_font = Font(color="9C0006")

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = copy.copy(header_fill)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for record in records:
        ws.append([record.get(h) for h in headers])
        row = ws.max_row

        # I列“京能物料号”按文本保存，避免长纯数字显示为科学计数法。
        write_material_no(ws.cell(row, 9), record.get("京能物料号"))
        status = str(record.get("匹配状态") or "")
        if status != "已匹配":
            for col in range(1, len(headers) + 1):
                ws.cell(row, col).fill = copy.copy(error_fill)
            ws.cell(row, 11).font = copy.copy(error_font)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{max(ws.max_row, 1)}"

    widths = {
        "A": 16, "B": 10, "C": 18, "D": 22, "E": 24, "F": 16,
        "G": 18, "H": 12, "I": 18, "J": 12, "K": 28, "L": 14,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    # 输出文件用 Excel 打开时，默认显示“京能物料对应汇总”。
    # 仅改变工作簿的活动 Sheet，不影响其它 Sheet 的内容和处理逻辑。
    for sheet in wb.worksheets:
        sheet.sheet_view.tabSelected = False
    ws.sheet_view.tabSelected = True
    wb.active = ws


# ============================================================
# 主映射表索引
# ============================================================


def extract_mapping_pairs(ws, row: int) -> Tuple[List[Tuple[Any, Decimal]], Optional[str]]:
    """
    从 M:X 读取物料A~F。

    返回：
      ([(我方物料号, 单位映射数量), ...], 错误信息)

    规则：
    - 物料号为空：该组忽略。
    - 物料号有值、数量为空：数量默认为 1。
    - 数量非空但不是数字：视为异常，不静默按 1。
    - 汇总表 G 列完全不参与数量计算。
    """
    start_col = MAPPING_START_COL
    pairs: List[Tuple[Any, Decimal]] = []

    for i in range(CONFIG["mapping_count"]):
        material_col = start_col + i * 2
        qty_col = material_col + 1
        material = ws.cell(row, material_col).value
        mapping_qty = ws.cell(row, qty_col).value

        if is_empty(material):
            continue

        if is_empty(mapping_qty):
            factor = Decimal("1")
        else:
            try:
                factor = to_decimal(mapping_qty)
            except ValueError:
                label = chr(ord("A") + i)
                return [], f"物料{label}数量异常：{mapping_qty}"
            assert factor is not None

        pairs.append((material, factor))

    return pairs, None


def build_master_indexes(master_path: Path):
    """
    一次流式扫描主映射表。
    read_only=True 配合 iter_rows(values_only=True)，避免逐单元格随机访问造成明显延迟。
    """
    if not master_path.exists():
        raise FileNotFoundError(f"主映射汇总表不存在：{master_path}")

    wb = load_workbook(master_path, read_only=True, data_only=True)
    try:
        ws = wb["汇总"] if "汇总" in wb.sheetnames else wb.active

        b_index: Dict[str, List[int]] = defaultdict(list)
        d_index: Dict[str, List[int]] = defaultdict(list)
        mappings: Dict[int, Tuple[List[Tuple[Any, Decimal]], Optional[str]]] = {}

        max_col = MAPPING_MAX_COL
        b_index_0 = col_to_num(CONFIG["master_material_col"]) - 1
        d_index_0 = col_to_num(CONFIG["master_model_col"]) - 1

        for row_no, values in enumerate(
            ws.iter_rows(
                min_row=2,
                max_col=max_col,
                values_only=True,
            ),
            start=2,
        ):
            b_value = values[b_index_0] if b_index_0 < len(values) else None
            d_value = values[d_index_0] if d_index_0 < len(values) else None

            b_key = normalize_key(b_value)
            d_key = normalize_key(d_value)

            if b_key:
                b_index[b_key].append(row_no)
            if d_key:
                d_index[d_key].append(row_no)

            mappings[row_no] = extract_mapping_pairs_from_values(values)

        common_records: List[Dict[str, Any]] = []
        common_sheet_name = CONFIG["common_material_sheet_name"]
        if common_sheet_name in wb.sheetnames:
            common_records = build_common_material_definitions(wb[common_sheet_name])

        return dict(b_index), dict(d_index), mappings, common_records
    finally:
        wb.close()


# ============================================================
# 客户行匹配
# ============================================================


def find_master_row(
    customer_b: Any,
    customer_d: Any,
    b_index: Dict[str, List[int]],
    d_index: Dict[str, List[int]],
) -> Tuple[Optional[int], str]:
    """
    严格匹配规则：
    - 客户 B 有值：只按 B 匹配；B 未找到时不再自动回退 D。
    - 客户 B 为空：才按 D 匹配。
    - 多条候选记录：视为映射冲突，不自动猜测。
    """
    b_key = normalize_key(customer_b)
    d_key = normalize_key(customer_d)

    if b_key:
        rows = b_index.get(b_key, [])
        if len(rows) == 1:
            return rows[0], ""
        if len(rows) == 0:
            return None, "未找到映射（按物料号B）"
        return None, f"映射冲突：物料号B匹配到{len(rows)}条"

    if not d_key:
        return None, "未找到映射：物料号B和型号D均为空"

    rows = d_index.get(d_key, [])
    if len(rows) == 1:
        return rows[0], ""
    if len(rows) == 0:
        return None, "未找到映射（按型号D）"
    return None, f"映射冲突：型号D匹配到{len(rows)}条"


def collect_source_rows(ws, sheet_cfg: Dict[str, Any]) -> List[int]:
    """按既有规则，从 start_row 开始，遇 D 为空即结束；按列顺序迭代减少随机单元格访问。"""
    rows: List[int] = []
    start_row = sheet_cfg["start_row"]
    stop_col = col_to_num(sheet_cfg["stop_col"])
    for row_no, (cell,) in enumerate(
        ws.iter_rows(min_row=start_row, max_row=ws.max_row, min_col=stop_col, max_col=stop_col),
        start=start_row,
    ):
        if is_empty(cell.value):
            break
        rows.append(row_no)
    return rows


def process_sheet(
    ws,
    sheet_cfg: Dict[str, Any],
    b_index: Dict[str, List[int]],
    d_index: Dict[str, List[int]],
    mappings: Dict[int, Tuple[List[Tuple[Any, Decimal]], Optional[str]]],
) -> Tuple[Dict[str, int], List[Dict[str, Any]]]:
    start_row = sheet_cfg["start_row"]
    header_row = start_row - 1
    material_col = col_to_num(sheet_cfg["material_col"])
    model_col = col_to_num(sheet_cfg["model_col"])
    quantity_col = col_to_num(sheet_cfg["quantity_col"])
    style_source_col = quantity_col
    apply_result_headers(ws, header_row, style_source_col)

    source_rows = collect_source_rows(ws, sheet_cfg)

    stats = {
        "source_rows": len(source_rows),
        "matched_rows": 0,
        "expanded_rows": 0,
        "unmatched_rows": 0,
        "conflict_rows": 0,
        "quantity_error_rows": 0,
        "mapping_error_rows": 0,
        "no_mapping_rows": 0,
    }
    summary_records: List[Dict[str, Any]] = []

    # 必须从下往上：插入行后不会改变尚未处理的原始行号。
    for row in reversed(source_rows):
        clear_result_cells(ws, row)

        customer_b = ws.cell(row, material_col).value
        customer_d = ws.cell(row, model_col).value
        customer_qty_raw = ws.cell(row, quantity_col).value

        master_row, match_error = find_master_row(customer_b, customer_d, b_index, d_index)
        if master_row is None:
            set_status(ws, row, match_error)
            highlight_error_row(ws, row)
            if match_error.startswith("映射冲突"):
                stats["conflict_rows"] += 1
            else:
                stats["unmatched_rows"] += 1
            summary_records.append(
                build_summary_record(ws, sheet_cfg, row, None, None, match_error)
            )
            continue

        try:
            customer_qty = to_decimal(customer_qty_raw)
        except ValueError:
            customer_qty = None

        if customer_qty is None:
            status = f"客户数量异常：{customer_qty_raw}"
            set_status(ws, row, status)
            highlight_error_row(ws, row)
            stats["quantity_error_rows"] += 1
            summary_records.append(
                build_summary_record(ws, sheet_cfg, row, None, None, status)
            )
            continue

        pairs, mapping_error = mappings[master_row]
        if mapping_error:
            status = f"映射数量异常：{mapping_error}"
            set_status(ws, row, status)
            highlight_error_row(ws, row)
            stats["mapping_error_rows"] += 1
            summary_records.append(
                build_summary_record(ws, sheet_cfg, row, None, None, status)
            )
            continue

        if not pairs:
            status = "未配置我方物料"
            set_status(ws, row, status)
            highlight_error_row(ws, row)
            stats["no_mapping_rows"] += 1
            summary_records.append(
                build_summary_record(ws, sheet_cfg, row, None, None, status)
            )
            continue

        # 1:N 展开。先在原行下方插入 N-1 行，再复制 A:Z。
        extra = len(pairs) - 1
        if extra > 0:
            ws.insert_rows(row + 1, amount=extra)
            for offset in range(1, extra + 1):
                copy_row_a_to_z(ws, row, row + offset)
            stats["expanded_rows"] += extra

        for offset, (our_material, factor) in enumerate(pairs):
            target_row = row + offset
            clear_result_cells(ws, target_row)
            result_qty = customer_qty * factor
            write_material_no(
                ws.cell(target_row, RESULT_MATERIAL_COL),
                our_material,
            )
            result_qty_excel = decimal_to_excel(result_qty)
            ws.cell(target_row, RESULT_QTY_COL).value = result_qty_excel
            set_status(ws, target_row, "已匹配")
            summary_records.append(
                build_summary_record(
                    ws, sheet_cfg, row, our_material, result_qty_excel, "已匹配"
                )
            )

        stats["matched_rows"] += 1

    # 由于处理顺序为自下而上，这里按原始行号恢复客户顺序。
    # sort 为稳定排序，因此同一原始行的 1:N 映射仍保持 A~F 顺序。
    summary_records.sort(key=lambda record: record["原始行号"])
    return stats, summary_records


# ============================================================
# .xls 转换（保留原始 .xls，不修改源文件）
# ============================================================


class ExcelXlsConverter:
    """一个项目共用一个 Excel COM 进程，避免每个 .xls 都重复启动/退出 Excel。"""

    def __init__(self) -> None:
        self.excel = None

    def __enter__(self):
        try:
            import win32com.client  # type: ignore
        except ImportError as exc:
            raise RuntimeError(
                "检测到 .xls 文件。处理 .xls 需要 pywin32 和本机 Microsoft Excel。"
            ) from exc
        self.excel = win32com.client.DispatchEx("Excel.Application")
        self.excel.Visible = False
        self.excel.DisplayAlerts = False
        self.excel.ScreenUpdating = False
        try:
            self.excel.EnableEvents = False
        except Exception:
            pass
        return self

    def convert(self, source: Path, temp_dir: Path) -> Path:
        if self.excel is None:
            raise RuntimeError("Excel 转换器尚未启动。")
        out = temp_dir / f"{source.stem}.xlsx"
        workbook = None
        try:
            workbook = self.excel.Workbooks.Open(
                str(source.resolve()),
                ReadOnly=True,
                UpdateLinks=0,
                AddToMru=False,
            )
            workbook.SaveAs(str(out.resolve()), FileFormat=51)  # xlOpenXMLWorkbook (.xlsx)
            workbook.Close(SaveChanges=False)
            workbook = None
            return out
        except Exception as exc:
            raise RuntimeError(f".xls 转换失败：{exc}") from exc
        finally:
            if workbook is not None:
                try:
                    workbook.Close(SaveChanges=False)
                except Exception:
                    pass

    def __exit__(self, exc_type, exc, tb) -> None:
        if self.excel is not None:
            try:
                self.excel.Quit()
            except Exception:
                pass
            self.excel = None


def convert_xls_to_xlsx_with_excel(source: Path, temp_dir: Path) -> Path:
    """兼容单文件调用；项目批处理时由 run_project 复用 ExcelXlsConverter。"""
    with ExcelXlsConverter() as converter:
        return converter.convert(source, temp_dir)


# ============================================================
# 单文件处理
# ============================================================


def build_output_name(source: Path) -> str:
    if source.suffix.lower() == ".xls":
        return f"{source.stem}_已匹配.xlsx"
    return f"{source.stem}_已匹配{source.suffix.lower()}"


def process_one_file(
    source: Path,
    project_dir: Path,
    b_index: Dict[str, List[int]],
    d_index: Dict[str, List[int]],
    mappings: Dict[int, Tuple[List[Tuple[Any, Decimal]], Optional[str]]],
    common_definitions: List[Dict[str, Any]],
    quote_cache: Dict[Tuple[str, str], Tuple[Path, List[Dict[str, Any]]]],
    xls_converter: Optional[ExcelXlsConverter] = None,
) -> Tuple[Path, Dict[str, Dict[str, int]], Dict[str, Any]]:
    suffix = source.suffix.lower()
    if suffix not in (".xls", ".xlsx", ".xlsm"):
        raise ValueError(f"不支持的文件类型：{suffix}")

    quote_info: Dict[str, Any] = {"records": 0, "abnormal": 0}
    common_records: List[Dict[str, Any]] = []
    if common_definitions:
        quote_file, quote_block = resolve_quote_block(source, project_dir, quote_cache)
        common_records, common_abnormal = build_common_records_from_quote(
            common_definitions, quote_block
        )
        quote_info = {
            "quote_file": str(quote_file),
            "quote_group": quote_block["group_text"],
            "quote_row": quote_block["start_row"],
            "records": len(common_records),
            "abnormal": common_abnormal,
        }

    temp_ctx = tempfile.TemporaryDirectory(prefix="material_fillback_")
    try:
        temp_dir = Path(temp_ctx.name)
        working_file = source
        keep_vba = suffix == ".xlsm"

        if suffix == ".xls":
            if xls_converter is not None:
                working_file = xls_converter.convert(source, temp_dir)
            else:
                working_file = convert_xls_to_xlsx_with_excel(source, temp_dir)
            keep_vba = False

        wb = load_workbook(
            working_file,
            data_only=False,
            keep_vba=keep_vba,
        )
        try:
            file_stats: Dict[str, Dict[str, int]] = {}
            summary_records: List[Dict[str, Any]] = []
            processed_sheet_count = 0

            for sheet_cfg in CONFIG["sheets"]:
                name = sheet_cfg["name"]
                if name not in wb.sheetnames:
                    log(f"    [跳过] Sheet 不存在：{name}")
                    continue

                stats, sheet_records = process_sheet(
                    wb[name], sheet_cfg, b_index, d_index, mappings
                )
                file_stats[name] = stats
                summary_records.extend(sheet_records)
                processed_sheet_count += 1

            if processed_sheet_count == 0:
                raise ValueError("未找到需要处理的 Sheet：电柜元器件 / 电柜半成品")

            # 统一追加主映射表中的公共物料。
            summary_records.extend(common_records)

            create_summary_sheet(wb, summary_records)

            output_path = unique_path(OUTPUT_DIR, build_output_name(source))
            wb.save(output_path)
            return output_path, file_stats, quote_info
        finally:
            wb.close()
    finally:
        temp_ctx.cleanup()


# ============================================================
# 主程序
# ============================================================


def run_project(project_dir: Path) -> Dict[str, Any]:
    """按项目目录执行脚本3。原始文件只读，结果写入 项目\输出\物料回填。"""
    from project_utils import iter_project_excel_files, project_log_dir, project_output_dir

    global OUTPUT_DIR, LOG_FILE
    _log_lines.clear()
    project_dir = Path(project_dir).resolve()
    OUTPUT_DIR = project_output_dir(project_dir) / "物料回填"
    LOG_FILE = project_log_dir(project_dir) / "03_生成京能物料对应表.log"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    LOG_FILE.parent.mkdir(parents=True, exist_ok=True)

    log("=" * 64)
    log("脚本3：生成客户清单京能物料对应表")
    log("=" * 64)
    log(f"项目目录：{project_dir}")
    log(f"主映射表：{MASTER_FILE}")
    log(f"输出目录：{OUTPUT_DIR}")
    log("原始项目文件只读，不移动。")
    log("")

    files = iter_project_excel_files(project_dir)
    if not files:
        result = {"success": False, "can_continue": False, "abnormal_rows": 0,
                  "reason": "项目目录中没有找到可处理的电气清单 Excel 文件。"}
        log(f"[错误] {result['reason']}")
        save_log()
        return result

    log(f"发现客户清单：{len(files)} 个")
    log("正在读取主映射汇总表...")
    try:
        b_index, d_index, mappings, common_definitions = build_master_indexes(MASTER_FILE)
    except Exception as exc:
        result = {"success": False, "can_continue": False, "abnormal_rows": 0,
                  "reason": f"无法读取主映射汇总表：{exc}"}
        log(f"[错误] {result['reason']}")
        save_log()
        return result

    log(f"公共物料配置：{len(common_definitions)} 条；数量将在每张电气清单对应的报价柜型块中提取。\n")
    ok_count = fail_count = 0
    total_source_rows = total_matched_rows = total_expanded_rows = 0
    total_abnormal_rows = 0
    total_common_records = total_common_abnormal = 0
    output_files: List[str] = []
    quote_cache: Dict[Tuple[str, str], Tuple[Path, List[Dict[str, Any]]]] = {}

    has_xls = any(path.suffix.lower() == ".xls" for path in files)
    converter_ctx = ExcelXlsConverter() if has_xls else None
    converter = None
    batch_started = time.perf_counter()

    try:
        if converter_ctx is not None:
            log("检测到 .xls：本项目只启动一次 Excel，批量完成格式转换。")
            try:
                converter = converter_ctx.__enter__()
            except Exception as exc:
                result = {
                    "success": False,
                    "can_continue": False,
                    "abnormal_rows": total_abnormal_rows,
                    "reason": f"无法启动 Excel 批量转换：{exc}",
                    "output_files": output_files,
                    "log_file": str(LOG_FILE),
                }
                log(f"[错误] {result['reason']}")
                save_log()
                return result

        for index, source in enumerate(files, start=1):
            file_started = time.perf_counter()
            log(f"[{index}/{len(files)}] {source.relative_to(project_dir)}")
            try:
                output_path, file_stats, quote_info = process_one_file(
                    source, project_dir, b_index, d_index, mappings,
                    common_definitions, quote_cache, converter
                )
                output_files.append(str(output_path))
                if quote_info.get("quote_file"):
                    log(
                        f"    报价：{Path(quote_info['quote_file']).name} / "
                        f"柜号 {quote_info.get('quote_group')} / 起始行 {quote_info.get('quote_row')}"
                    )
                    log(
                        f"    公共物料：展开{quote_info.get('records', 0)}条，"
                        f"异常{quote_info.get('abnormal', 0)}条"
                    )
                    total_common_records += int(quote_info.get("records", 0))
                    total_common_abnormal += int(quote_info.get("abnormal", 0))
                    total_abnormal_rows += int(quote_info.get("abnormal", 0))
                for sheet_name, stats in file_stats.items():
                    abnormal = (stats["unmatched_rows"] + stats["conflict_rows"] +
                                stats["quantity_error_rows"] + stats["mapping_error_rows"] +
                                stats["no_mapping_rows"])
                    total_source_rows += stats["source_rows"]
                    total_matched_rows += stats["matched_rows"]
                    total_expanded_rows += stats["expanded_rows"]
                    total_abnormal_rows += abnormal
                    log(f"    {sheet_name}: 原始{stats['source_rows']}行，匹配{stats['matched_rows']}行，展开新增{stats['expanded_rows']}行，异常{abnormal}行")
                ok_count += 1
                log(f"    [完成] 输出：{output_path.name}，耗时 {time.perf_counter() - file_started:.2f}s")
            except Exception as exc:
                fail_count += 1
                log(f"    [失败] {exc}，耗时 {time.perf_counter() - file_started:.2f}s")
            log("")
    finally:
        if converter_ctx is not None and converter is not None:
            converter_ctx.__exit__(None, None, None)

    total_elapsed = time.perf_counter() - batch_started

    log("=" * 64)
    log(f"文件总数：{len(files)}，成功：{ok_count}，失败：{fail_count}")
    log(f"客户原始物料行：{total_source_rows}")
    log(f"成功匹配客户行：{total_matched_rows}")
    log(f"1:N 展开新增行：{total_expanded_rows}")
    log(f"公共物料输出记录：{total_common_records}")
    log(f"公共物料异常：{total_common_abnormal}")
    log(f"总异常记录：{total_abnormal_rows}")
    log(f"脚本3总耗时：{total_elapsed:.2f}s")
    log(f"日志：{LOG_FILE}")
    save_log()

    if fail_count:
        return {"success": False, "can_continue": False, "abnormal_rows": total_abnormal_rows,
                "reason": f"有 {fail_count} 个文件处理失败，请查看项目日志。",
                "output_files": output_files, "log_file": str(LOG_FILE)}
    if total_abnormal_rows:
        return {"success": True, "can_continue": False, "abnormal_rows": total_abnormal_rows,
                "reason": f"物料回填存在 {total_abnormal_rows} 条异常，未执行 ERP 生成。",
                "output_files": output_files, "log_file": str(LOG_FILE)}
    return {"success": True, "can_continue": True, "abnormal_rows": 0, "reason": "",
            "output_files": output_files, "files": len(files), "log_file": str(LOG_FILE)}


def main() -> None:
    import argparse
    from project_utils import add_project_args, ensure_project_dir, write_result_json
    parser = argparse.ArgumentParser(description="脚本3：生成京能物料对应表")
    add_project_args(parser)
    args = parser.parse_args()
    try:
        result = run_project(ensure_project_dir(args.project))
    except Exception as exc:
        result = {"success": False, "can_continue": False, "abnormal_rows": 0, "reason": str(exc)}
        print(f"[错误] {exc}")
    write_result_json(args.result_json, result)
    raise SystemExit(0 if result.get("success") else 1)


if __name__ == "__main__":
    main()
